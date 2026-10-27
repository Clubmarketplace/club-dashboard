"""
Tela "Reversos" do gestor (27/10) -- visão completa do que voltou para o galpão:
quantos chegaram, quantos voltaram ao estoque, quantos não, o que está parado,
quem recebeu, quem concluiu e o tempo médio de tratamento.

  GET /api/reversos-gestor/resumo     ?dias=1|7|15|30  (ou ?de=&ate=)  &galpao=  -> JSON
  GET /api/reversos-gestor/card                                                   -> JSON (card do dashboard)
  GET /api/reversos-gestor/relatorio  ?dias=... (ou ?de=&ate=)  &galpao=          -> Excel (.xlsx)
  GET /api/reversos-gestor/lista      ?grupo=&situacao=&dia=&conta=&pagina=...    -> JSON (v22: clique nos cartões)
  GET /api/reversos-gestor/lista/excel  (mesmos filtros, sem paginação)           -> Excel (.xlsx)

Regras:
  - SÓ admin e supervisor (decisão do Washington, 27/10). Os demais perfis recebem 403,
    inclusive gerente de setor e logística. A página (/reversos-gestao) tem a mesma trava.
  - Entram TODOS os reversos: os registrados pelo galpão e os avisados pelo próprio seller.
    "Quem recebeu" lista só quem registrou no galpão (origem "logistica").
  - FOTOS NUNCA são lidas aqui (nem o conteúdo, nem a coluna legada): só números e textos.
    Foi a causa da queda de 26/10; a tela é de gestão e não precisa delas.
  - Datas em horário de Brasília; o banco guarda UTC cru. O "dia" é o do registro (criado_em).
  - "Fila agora" (aguardando / parados) NÃO depende do período: é o que está pendente hoje.
"""
import io
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta, timezone
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, Request
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from sqlalchemy import func
from sqlalchemy.orm import Session, defer

from app import auth
from app.database import get_db
from app.models import CustoSku, SolicitacaoCancelamento
from app.routers.painel_galpao import (
    _dia_br,
    _nome_situacao,
    _pct,
    _situacao_reverso,
    _tempo_medio_min,
    formatar_tempo,
)
from app.routers.solicitacoes_cancelamento import (
    FUSO_BR,
    GALPOES,
    PLATAFORMAS,
    TIPO_REVERSO,
    _limites_utc_produtividade,
    _texto_protocolo,
)

router = APIRouter(prefix="/api/reversos-gestor", tags=["reversos-gestor"])

PAPEIS_GESTOR_REVERSOS = ("admin", "supervisor")
PERIODOS_VALIDOS = (1, 7, 15, 30)   # "Hoje", 7, 15 e 30 dias (+ De/Até livre)
MAX_DIAS_INTERVALO = 366            # no máximo ~1 ano por consulta
DIAS_PARADO = 3                     # reverso aguardando há mais que isso = "parado" (alerta)
TOP_CONTAS = 15


# ---------------------------------------------------------------------------
# Acesso e período
# ---------------------------------------------------------------------------
def usuario_gestor_reversos(request: Request, db: Session):
    """Só admin e supervisor. 401 sem sessão, 403 para qualquer outro perfil."""
    usuario = auth.usuario_atual(request, db)
    if usuario is None:
        raise HTTPException(status_code=401, detail="Sessão expirada -- faça login de novo.")
    if usuario.papel not in PAPEIS_GESTOR_REVERSOS:
        raise HTTPException(status_code=403, detail="A tela de reversos é só para admin e supervisor.")
    return usuario


def _intervalo(dias: Optional[int], de: Optional[str], ate: Optional[str]) -> tuple[date, date]:
    """De/Até (livre) tem prioridade; senão usa 'dias' (1 = hoje). Valida tudo com 422."""
    hoje = datetime.now(FUSO_BR).date()
    if de or ate:
        try:
            data_de = date.fromisoformat(de) if de else hoje
            data_ate = date.fromisoformat(ate) if ate else hoje
        except ValueError:
            raise HTTPException(status_code=422, detail="Datas no formato AAAA-MM-DD.")
        if data_ate < data_de:
            raise HTTPException(status_code=422, detail="A data final não pode ser antes da inicial.")
        if (data_ate - data_de).days > MAX_DIAS_INTERVALO:
            raise HTTPException(status_code=422, detail="Escolha no máximo 1 ano de intervalo.")
        return data_de, data_ate
    dias = 7 if dias is None else dias
    if dias not in PERIODOS_VALIDOS:
        raise HTTPException(status_code=422, detail="Período inválido. Use 1 (hoje), 7, 15 ou 30 dias, ou De/Até.")
    return hoje - timedelta(days=dias - 1), hoje


def _checar_galpao(galpao: Optional[int]) -> None:
    if galpao is not None and galpao not in GALPOES:
        raise HTTPException(status_code=422, detail="Galpão inválido.")


# ---------------------------------------------------------------------------
# Consultas (sem fotos)
# ---------------------------------------------------------------------------
def _base(db: Session):
    return (
        db.query(SolicitacaoCancelamento)
        # Fotos (base64 legado) NÃO são lidas: tela só de números (ver cabeçalho).
        .options(defer(SolicitacaoCancelamento.fotos_base64))
        .filter(SolicitacaoCancelamento.tipo == TIPO_REVERSO)
        .filter(SolicitacaoCancelamento.excluido_em.is_(None))
    )


def _reversos_do_periodo(db: Session, de: date, ate: date, galpao: Optional[int]):
    inicio, fim = _limites_utc_produtividade(de, ate)
    consulta = _base(db).filter(SolicitacaoCancelamento.criado_em >= inicio, SolicitacaoCancelamento.criado_em < fim)
    if galpao is not None:
        consulta = consulta.filter(SolicitacaoCancelamento.galpao == galpao)
    return consulta.order_by(SolicitacaoCancelamento.criado_em.asc()).all()


def _reversos_aguardando(db: Session, galpao: Optional[int]):
    """Fila agora: tudo que ainda espera decisão da equipe, de qualquer data."""
    consulta = _base(db).filter(SolicitacaoCancelamento.confirmado_por.is_(None))
    if galpao is not None:
        consulta = consulta.filter(SolicitacaoCancelamento.galpao == galpao)
    return consulta.order_by(SolicitacaoCancelamento.criado_em.asc()).all()


# ---------------------------------------------------------------------------
# Cálculos (funções puras -- testáveis sem banco)
# ---------------------------------------------------------------------------
def _idade_min(s: SolicitacaoCancelamento, agora_utc: datetime) -> int:
    return max(int((agora_utc - s.criado_em).total_seconds() // 60), 0) if s.criado_em else 0


def montar_fila(aguardando, agora_utc: Optional[datetime] = None) -> dict:
    """Reversos pendentes agora: quantos, quantos 'parados' (> DIAS_PARADO dias) e o mais antigo."""
    agora_utc = agora_utc or datetime.utcnow()
    idades = [_idade_min(s, agora_utc) for s in aguardando]
    limite_min = DIAS_PARADO * 1440
    return {
        "aguardando": len(aguardando),
        "parados": sum(1 for i in idades if i > limite_min),
        "dias_parado": DIAS_PARADO,
        "mais_antigo_min": max(idades) if idades else None,
    }


TOP_SKUS = 10
SEM_SKU = "__sem__"   # valor enviado pela tela para "Sem SKU informado"
# 27/10: GARANTIA é um destino próprio (a equipe emite a nota). Nos cartões, no gráfico por dia e no
# ranking ela aparece separada de "não voltaram". Nas tabelas por galpão/conta e no Excel,
# "Fora do estoque" = não voltou + garantia (tudo que não foi para o estoque).
GARANTIA = "garantia"


def _chave_sku(s: SolicitacaoCancelamento) -> str:
    """SKU do reverso: a coluna sku (lida do pedido no ML) e, na falta dela, o campo 'produto'
    (produto_titulo), que no reverso é o SKU digitado. Reversos antigos podem estar sem nenhum
    dos dois = SEM_SKU. Maiúsculo para 'mpr530' e 'MPR530' juntarem."""
    return (s.sku or s.produto_titulo or "").strip().upper() or SEM_SKU


def montar_ranking_sku(linhas) -> list[dict]:
    """Top SKUs com mais reversos no período (função pura): recebidos e o destino de cada um."""
    grupos = defaultdict(list)
    for s in linhas:
        grupos[_chave_sku(s)].append(s)
    ranking = []
    for k, rs in grupos.items():
        c = Counter(_situacao_reverso(s) for s in rs)
        v, a, g = c.get("voltou", 0), c.get("aguardando", 0), c.get(GARANTIA, 0)
        ranking.append({"sku": k, "nome": "Sem SKU informado" if k == SEM_SKU else k,
                        "recebidos": len(rs), "voltaram": v, "aguardando": a, "garantias": g,
                        "nao_voltaram": len(rs) - v - a - g})
    ranking.sort(key=lambda x: (x["sku"] == SEM_SKU, -x["recebidos"], -x["nao_voltaram"], x["sku"]))
    return ranking


def montar_comparacao(atual: dict, anterior_linhas, de_ant: date, ate_ant: date) -> dict:
    """Período anterior (mesmo tamanho, logo antes) para as setas dos cartões. Função pura."""
    c = Counter(_situacao_reverso(s) for s in anterior_linhas)
    v, a, g = c.get("voltou", 0), c.get("aguardando", 0), c.get(GARANTIA, 0)
    return {"de": de_ant.isoformat(), "ate": ate_ant.isoformat(), "recebidos": len(anterior_linhas),
            "voltaram": v, "garantias": g, "nao_voltaram": len(anterior_linhas) - v - a - g}


def montar_resumo_gestor(linhas, de: date, ate: date, aguardando_agora, agora_utc: Optional[datetime] = None) -> dict:
    """Números da tela a partir dos reversos do período + a fila de pendentes de agora."""
    sit = Counter(_situacao_reverso(s) for s in linhas)
    voltaram = sit.get("voltou", 0)
    aguardando = sit.get("aguardando", 0)
    garantias = sit.get(GARANTIA, 0)
    nao_voltaram = len(linhas) - voltaram - aguardando - garantias
    concluidos = voltaram + nao_voltaram + garantias
    avisados_seller = sum(1 for s in linhas if s.origem != "logistica")

    # Por dia (todos os dias do período, mesmo vazios), com a situação de cada um
    por_dia = {de + timedelta(days=i): {"recebidos": 0, "voltou": 0, "garantia": 0, "nao_voltou": 0, "aguardando": 0}
               for i in range((ate - de).days + 1)}
    for s in linhas:
        d = _dia_br(s)
        if d in por_dia:
            sit_ = _situacao_reverso(s)
            por_dia[d]["recebidos"] += 1
            por_dia[d]["voltou" if sit_ == "voltou" else "aguardando" if sit_ == "aguardando" else "garantia" if sit_ == GARANTIA else "nao_voltou"] += 1

    # Resultado (rosca) e motivos de "não voltou ao estoque"
    ordem = ["voltou", GARANTIA, "avariado", "faltando_peca", "diferente", "negociacao_ml", "descartado", "outro", "aguardando"]
    extras = [k for k in sit if k not in ordem]       # legados (negociacao_ressarcido etc.)
    resultado = [{"chave": k, "nome": _nome_situacao(k), "quantidade": sit[k]} for k in ordem + extras if sit.get(k)]

    # Por galpão (os 3 sempre, + "sem galpão" quando o seller avisou) e total
    def bloco(rs):
        c = Counter(_situacao_reverso(s) for s in rs)
        v, a = c.get("voltou", 0), c.get("aguardando", 0)
        n = len(rs) - v - a            # "fora do estoque" = não voltou + garantia
        return {"recebidos": len(rs), "voltaram": v, "nao_voltaram": n, "aguardando": a, "garantias": c.get(GARANTIA, 0),
                "pct_voltaram": _pct(v, v + n), "tempo_medio_min": _tempo_medio_min(rs)}

    por_galpao = [{"galpao": n, "nome": nome, **bloco([s for s in linhas if s.galpao == n])} for n, nome in GALPOES.items()]
    sem_galpao = [s for s in linhas if s.galpao not in GALPOES]
    if sem_galpao:
        por_galpao.append({"galpao": None, "nome": "Sem galpão (aviso do seller)", **bloco(sem_galpao)})
    total = bloco(linhas)

    # Quem recebeu (registrou no galpão) e quem concluiu (equipe que decidiu)
    receb = defaultdict(int)
    for s in linhas:
        if s.origem == "logistica":
            receb[(s.solicitado_por or "Não informado", s.galpao)] += 1
    quem_recebeu = sorted(
        ({"nome": n, "galpao": g, "galpao_nome": GALPOES.get(g, "—"), "quantidade": q} for (n, g), q in receb.items()),
        key=lambda x: (-x["quantidade"], x["nome"].lower()),
    )
    por_pessoa = defaultdict(list)
    for s in linhas:
        if s.confirmado_por:
            por_pessoa[s.confirmado_por].append(s)
    quem_concluiu = sorted(
        ({"nome": nome, "concluidos": len(rs),
          "voltaram": sum(1 for s in rs if _situacao_reverso(s) == "voltou"),
          "nao_voltaram": sum(1 for s in rs if _situacao_reverso(s) != "voltou"),
          "tempo_medio_min": _tempo_medio_min(rs)} for nome, rs in por_pessoa.items()),
        key=lambda x: (-x["concluidos"], x["nome"].lower()),
    )

    # Por conta (quais contas mais têm reverso e o que aconteceu com eles)
    por_conta = defaultdict(list)
    nomes = {}
    for s in linhas:
        k = s.conta_chave or (s.conta or "").lower()
        nomes.setdefault(k, s.conta)
        por_conta[k].append(s)
    contas = sorted(({"conta": nomes[k], **bloco(rs)} for k, rs in por_conta.items()),
                    key=lambda x: (-x["recebidos"], x["conta"].lower()))

    tempos = [max((s.confirmado_em - s.criado_em).total_seconds() / 60, 0)
              for s in linhas if s.confirmado_por and s.confirmado_em and s.criado_em]
    return {
        "periodo": {"de": de.isoformat(), "ate": ate.isoformat(), "dias": (ate - de).days + 1},
        "kpis": {
            **total, "avisados_seller": avisados_seller,
            "fora_estoque": total["nao_voltaram"],     # não voltou + garantia (as tabelas usam este)
            "nao_voltaram": nao_voltaram,              # só os que NÃO voltaram e não são garantia (cartão)
            "tempo_max_min": round(max(tempos)) if tempos else None,
            "conferidos": concluidos,
        },
        "fila": montar_fila(aguardando_agora, agora_utc),
        "por_dia": [{"data": d.isoformat(), **v} for d, v in sorted(por_dia.items())],
        "resultado": resultado,
        "por_galpao": por_galpao,
        "total_galpoes": total,
        "quem_recebeu": quem_recebeu,
        "quem_concluiu": quem_concluiu,
        "contas": contas[:TOP_CONTAS],
        "total_contas": len(contas),
        "ranking_sku": montar_ranking_sku(linhas)[:TOP_SKUS],
        "total_skus": len({_chave_sku(s) for s in linhas}),
    }


# ---------------------------------------------------------------------------
# Rotas
# ---------------------------------------------------------------------------
@router.get("/resumo")
def resumo(request: Request, dias: Optional[int] = None, de: Optional[str] = None, ate: Optional[str] = None,
           galpao: Optional[int] = None, db: Session = Depends(get_db)):
    usuario_gestor_reversos(request, db)
    _checar_galpao(galpao)
    d1, d2 = _intervalo(dias, de, ate)
    r = montar_resumo_gestor(_reversos_do_periodo(db, d1, d2, galpao), d1, d2, _reversos_aguardando(db, galpao))
    # Período anterior: mesmo tamanho, imediatamente antes (ex.: 7 dias -> os 7 dias antes). Só para as setas.
    tam = (d2 - d1).days + 1
    a2, a1 = d1 - timedelta(days=1), d1 - timedelta(days=tam)
    r["anterior"] = montar_comparacao(r, _reversos_do_periodo(db, a1, a2, galpao), a1, a2)
    return r


@router.get("/card")
def card(request: Request, db: Session = Depends(get_db)):
    """Resumo de HOJE + fila agora, para o card do dashboard do admin/supervisor."""
    usuario_gestor_reversos(request, db)
    hoje = datetime.now(FUSO_BR).date()
    r = montar_resumo_gestor(_reversos_do_periodo(db, hoje, hoje, None), hoje, hoje, _reversos_aguardando(db, None))
    k = r["kpis"]
    return {"hoje": {"recebidos": k["recebidos"], "voltaram": k["voltaram"], "nao_voltaram": k["nao_voltaram"]},
            "fila": r["fila"], "tempo_medio_min": k["tempo_medio_min"]}


def _hora_br(data_utc: Optional[datetime]) -> str:
    return data_utc.replace(tzinfo=timezone.utc).astimezone(FUSO_BR).strftime("%d/%m/%Y %H:%M") if data_utc else ""


@router.get("/relatorio")
def relatorio(request: Request, dias: Optional[int] = None, de: Optional[str] = None, ate: Optional[str] = None,
              galpao: Optional[int] = None, db: Session = Depends(get_db)):
    """Excel: Resumo, Por conta, Pessoas e Detalhe (uma linha por reverso, sem fotos)."""
    usuario = usuario_gestor_reversos(request, db)
    _checar_galpao(galpao)
    d1, d2 = _intervalo(dias, de, ate)
    linhas = _reversos_do_periodo(db, d1, d2, galpao)
    r = montar_resumo_gestor(linhas, d1, d2, _reversos_aguardando(db, galpao))
    k, f = r["kpis"], r["fila"]
    galpao_txt = GALPOES[galpao] if galpao else "Todos os galpões"

    wb = Workbook()
    aba = wb.active
    aba.title = "Resumo"
    cab_g = ["", "Recebidos", "No estoque", "Fora do estoque", "Aguardando", "% no estoque", "Tempo médio"]
    for linha in (
        ["Relatório de reversos"],
        ["Período", f"{d1.strftime('%d/%m/%Y')} a {d2.strftime('%d/%m/%Y')}"], ["Galpão", galpao_txt],
        ["Gerado em", f"{datetime.now(FUSO_BR).strftime('%d/%m/%Y %H:%M')} por {usuario.nome_exibicao}"], [],
        ["REVERSOS NO PERÍODO"], ["Recebidos", k["recebidos"]], ["Voltaram ao estoque", k["voltaram"]],
        ["Foram para garantia", k["garantias"]],
        ["Não voltaram ao estoque", k["nao_voltaram"]], ["Aguardando a equipe", k["aguardando"]],
        ["% que voltou ao estoque (dos conferidos)", f"{k['pct_voltaram']}%"],
        ["Tempo médio até a equipe decidir", formatar_tempo(k["tempo_medio_min"])],
        ["Avisados pelo próprio seller", k["avisados_seller"]], [],
        ["FILA AGORA (independe do período)"], ["Aguardando a equipe", f["aguardando"]],
        [f"Parados há mais de {f['dias_parado']} dias", f["parados"]],
        ["Mais antigo na fila", formatar_tempo(f["mais_antigo_min"])], [],
        ["RESULTADO DOS REVERSOS"], *([x["nome"], x["quantidade"]] for x in r["resultado"]), [],
        ["POR GALPÃO", *cab_g[1:]],
        *([g["nome"], g["recebidos"], g["voltaram"], g["nao_voltaram"], g["aguardando"], f"{g['pct_voltaram']}%",
           formatar_tempo(g["tempo_medio_min"])] for g in r["por_galpao"]),
        (lambda g: ["TOTAL", g["recebidos"], g["voltaram"], g["nao_voltaram"], g["aguardando"], f"{g['pct_voltaram']}%",
                    formatar_tempo(g["tempo_medio_min"])])(r["total_galpoes"]),
    ):
        aba.append(linha)
    aba.column_dimensions["A"].width = 42
    for col in "BCDEFG":
        aba.column_dimensions[col].width = 18

    aba_c = wb.create_sheet("Por conta")
    aba_c.append(["Conta", "Recebidos", "No estoque", "Fora do estoque", "Aguardando", "% no estoque", "Tempo médio"])
    por_conta = defaultdict(list)
    nomes = {}
    for s in linhas:   # reagrupa SEM o corte do top da tela
        k_ = s.conta_chave or (s.conta or "").lower()
        nomes.setdefault(k_, s.conta)
        por_conta[k_].append(s)
    for k_, rs in sorted(por_conta.items(), key=lambda kv: (-len(kv[1]), nomes[kv[0]].lower())):
        c = Counter(_situacao_reverso(s) for s in rs)
        v, a = c.get("voltou", 0), c.get("aguardando", 0)
        n = len(rs) - v - a
        aba_c.append([nomes[k_], len(rs), v, n, a, f"{_pct(v, v + n)}%", formatar_tempo(_tempo_medio_min(rs))])
    for col, larg in zip("ABCDEFG", (30, 12, 12, 16, 12, 14, 14)):
        aba_c.column_dimensions[col].width = larg

    aba_p = wb.create_sheet("Pessoas")
    aba_p.append(["QUEM RECEBEU NO GALPÃO", "Galpão", "Reversos recebidos"])
    for x in r["quem_recebeu"]:
        aba_p.append([x["nome"], x["galpao_nome"], x["quantidade"]])
    aba_p.append([])
    aba_p.append(["QUEM CONCLUIU (EQUIPE)", "Concluídos", "No estoque", "Fora do estoque", "Tempo médio"])
    for x in r["quem_concluiu"]:
        aba_p.append([x["nome"], x["concluidos"], x["voltaram"], x["nao_voltaram"], formatar_tempo(x["tempo_medio_min"])])
    for col, larg in zip("ABCDE", (32, 14, 14, 16, 14)):
        aba_p.column_dimensions[col].width = larg

    aba_d = wb.create_sheet("Detalhe")
    aba_d.append(["Recebido em", "Galpão", "Conta", "Plataforma", "Nº da venda", "Produto", "Aspecto / motivo", "Situação",
                  "Concluído por", "Concluído em", "Tempo até decidir", "Protocolo", "Observação da equipe", "Registrado por", "Origem"])
    for s in linhas:
        espera = None
        if s.confirmado_por and s.confirmado_em and s.criado_em:
            espera = max(round((s.confirmado_em - s.criado_em).total_seconds() / 60), 0)
        aba_d.append([
            _hora_br(s.criado_em), GALPOES.get(s.galpao, ""), s.conta, PLATAFORMAS.get(s.plataforma, s.plataforma), s.numero_venda,
            s.produto_titulo or s.sku or "", s.motivo or "", _nome_situacao(_situacao_reverso(s)), s.confirmado_por or "",
            _hora_br(s.confirmado_em), formatar_tempo(espera) if espera is not None else "", _texto_protocolo(s),
            s.observacao_resultado or "", s.solicitado_por or "", "Galpão" if s.origem == "logistica" else "Seller",
        ])
    for i, larg in enumerate((17, 11, 22, 14, 20, 36, 30, 32, 18, 17, 14, 18, 30, 18, 10), start=1):
        aba_d.column_dimensions[aba_d.cell(row=1, column=i).column_letter].width = larg

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    nome = f"reversos_{d1.isoformat()}_a_{d2.isoformat()}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{nome}"'},
    )


# ---------------------------------------------------------------------------
# Lista detalhada de um grupo (v22, 27/10): o gestor clica num cartão, numa barra
# do gráfico, numa fatia da rosca, num motivo ou numa conta e vê QUAIS reversos são,
# com conta, motivo, dia, quem recebeu, quem decidiu e quanto tempo levou.
# Continua SEM fotos (só texto e números), pelo mesmo motivo da queda de 26/10.
# ---------------------------------------------------------------------------
GRUPOS_LISTA = {
    "recebidos": "Todos os recebidos no período",
    "voltaram": "Voltaram ao estoque",
    "nao_voltaram": "Não voltaram ao estoque",
    "garantias": "Foram para garantia",
    "aguardando": "Aguardando a equipe (recebidos no período)",
    "fila": "Fila agora: aguardando a equipe",
    "parados": "Parados na fila há mais de %d dias" % DIAS_PARADO,
}
GRUPOS_SEM_PERIODO = ("fila", "parados")   # são o que está pendente HOJE, de qualquer data
POR_PAGINA_PADRAO = 50
POR_PAGINA_MAX = 200
TOP_RESUMO_LISTA = 8


def _espera_min(s: SolicitacaoCancelamento) -> Optional[int]:
    """Minutos entre o recebimento e a decisão da equipe (None se ainda não decidiu).
    Defensivo: data faltando ou fuso misturado não derruba a lista."""
    if s.confirmado_por and s.confirmado_em and s.criado_em:
        try:
            return max(round((s.confirmado_em - s.criado_em).total_seconds() / 60), 0)
        except TypeError:
            return None
    return None


def _item_lista(s: SolicitacaoCancelamento, agora_utc: datetime) -> dict:
    sit = _situacao_reverso(s)
    return {
        "id": s.id,
        "recebido_em": _hora_br(s.criado_em),
        "dia": _dia_br(s).isoformat() if s.criado_em else None,
        "galpao_nome": GALPOES.get(s.galpao, "Sem galpão"),
        "conta": s.conta or "—",
        "plataforma": PLATAFORMAS.get(s.plataforma, s.plataforma or ""),
        "numero_venda": s.numero_venda or "",
        "produto": s.produto_titulo or s.sku or "",
        "sku": s.sku or "",
        "aspecto": s.motivo or "",                       # o que o galpão/seller relatou ao receber
        "situacao": sit,
        "situacao_nome": _nome_situacao(sit),
        "observacao": s.observacao_resultado or "",
        "concluido_por": s.confirmado_por or "",
        "concluido_em": _hora_br(s.confirmado_em),
        "espera_min": _espera_min(s),                    # recebimento -> decisão
        "parado_ha_min": _idade_min(s, agora_utc) if sit == "aguardando" else None,
        "registrado_por": s.solicitado_por or "",
        "origem": "Galpão" if s.origem == "logistica" else "Seller",
        "protocolo": _texto_protocolo(s),
    }


def _selecionar_lista(db: Session, grupo: str, situacao: Optional[str], dia: Optional[str], conta: Optional[str],
                      dias: Optional[int], de: Optional[str], ate: Optional[str], galpao: Optional[int],
                      agora_utc: Optional[datetime] = None, sku: Optional[str] = None):
    """Devolve (linhas, d1, d2). d1/d2 = None nos grupos que não dependem do período."""
    if grupo not in GRUPOS_LISTA:
        raise HTTPException(status_code=422, detail="Grupo inválido.")
    _checar_galpao(galpao)
    dia_d = None
    if dia:
        try:
            dia_d = date.fromisoformat(dia)
        except ValueError:
            raise HTTPException(status_code=422, detail="Dia no formato AAAA-MM-DD.")
    agora_utc = agora_utc or datetime.utcnow()

    d1 = d2 = None
    if grupo in GRUPOS_SEM_PERIODO:
        linhas = _reversos_aguardando(db, galpao)
        if grupo == "parados":
            limite = DIAS_PARADO * 1440
            linhas = [s for s in linhas if _idade_min(s, agora_utc) > limite]
    else:
        d1, d2 = _intervalo(dias, de, ate)
        linhas = _reversos_do_periodo(db, d1, d2, galpao)
        if grupo == "voltaram":
            linhas = [s for s in linhas if _situacao_reverso(s) == "voltou"]
        elif grupo == "nao_voltaram":
            linhas = [s for s in linhas if _situacao_reverso(s) not in ("voltou", "aguardando", GARANTIA)]
        elif grupo == "garantias":
            linhas = [s for s in linhas if _situacao_reverso(s) == GARANTIA]
        elif grupo == "aguardando":
            linhas = [s for s in linhas if _situacao_reverso(s) == "aguardando"]

    if situacao:
        linhas = [s for s in linhas if _situacao_reverso(s) == situacao]
    if dia_d:
        linhas = [s for s in linhas if s.criado_em and _dia_br(s) == dia_d]
    if conta:
        q = conta.strip().lower()
        linhas = [s for s in linhas if (s.conta_chave or "").lower() == q or (s.conta or "").lower() == q]

    if sku:
        alvo = SEM_SKU if sku.strip().lower() == SEM_SKU else sku.strip().upper()
        linhas = [s for s in linhas if _chave_sku(s) == alvo]

    # Fila: o que espera há mais tempo primeiro. Demais: o mais recente primeiro.
    linhas = sorted(linhas, key=lambda s: s.criado_em or datetime.min, reverse=grupo not in GRUPOS_SEM_PERIODO)
    return linhas, d1, d2


def resumo_da_lista(linhas, agora_utc: Optional[datetime] = None) -> dict:
    """Resumo do que está na lista (função pura): motivos, contas, dias, galpões e quem decidiu."""
    agora_utc = agora_utc or datetime.utcnow()
    sit = Counter(_situacao_reverso(s) for s in linhas)
    voltaram, aguardando = sit.get("voltou", 0), sit.get("aguardando", 0)

    ordem = ["voltou", GARANTIA, "avariado", "faltando_peca", "diferente", "negociacao_ml", "descartado", "outro", "aguardando"]
    chaves = ordem + [k for k in sit if k not in ordem]
    por_motivo = [{"chave": k, "nome": _nome_situacao(k), "quantidade": sit[k]} for k in chaves if sit.get(k)]

    por_conta = defaultdict(list)
    nomes = {}
    for s in linhas:
        k = s.conta_chave or (s.conta or "").lower()
        nomes.setdefault(k, s.conta or "—")
        por_conta[k].append(s)
    contas = sorted(
        ({"conta": nomes[k], "quantidade": len(rs),
          "voltaram": sum(1 for s in rs if _situacao_reverso(s) == "voltou"),
          "aguardando": sum(1 for s in rs if _situacao_reverso(s) == "aguardando"),
          "nao_voltaram": sum(1 for s in rs if _situacao_reverso(s) not in ("voltou", "aguardando"))}
         for k, rs in por_conta.items()),
        key=lambda x: (-x["quantidade"], x["conta"].lower()))

    por_dia = Counter(_dia_br(s) for s in linhas if s.criado_em)
    dias_lista = [{"data": d.isoformat(), "quantidade": q} for d, q in sorted(por_dia.items(), reverse=True)]

    por_galpao = Counter(GALPOES.get(s.galpao, "Sem galpão") for s in linhas)
    decidiu = Counter(s.confirmado_por for s in linhas if s.confirmado_por)
    idades = [_idade_min(s, agora_utc) for s in linhas if _situacao_reverso(s) == "aguardando"]
    return {
        "total": len(linhas), "voltaram": voltaram, "aguardando": aguardando,
        "nao_voltaram": len(linhas) - voltaram - aguardando,
        "tempo_medio_min": _tempo_medio_min(linhas),
        "mais_antigo_min": max(idades) if idades else None,
        "avisados_seller": sum(1 for s in linhas if s.origem != "logistica"),
        "por_motivo": por_motivo,
        "contas": contas[:TOP_RESUMO_LISTA], "total_contas": len(contas),
        "dias": dias_lista[:31], "total_dias": len(dias_lista),
        "por_galpao": [{"nome": n, "quantidade": q} for n, q in por_galpao.most_common()],
        "quem_decidiu": [{"nome": n, "quantidade": q} for n, q in decidiu.most_common(TOP_RESUMO_LISTA)],
    }


@router.get("/lista")
def lista(request: Request, grupo: str = "recebidos", situacao: Optional[str] = None, dia: Optional[str] = None,
          conta: Optional[str] = None, dias: Optional[int] = None, de: Optional[str] = None, ate: Optional[str] = None,
          galpao: Optional[int] = None, pagina: int = 1, por_pagina: int = POR_PAGINA_PADRAO,
          sku: Optional[str] = None, db: Session = Depends(get_db)):
    """Lista paginada dos reversos de um grupo + resumo do grupo inteiro (não só da página)."""
    usuario_gestor_reversos(request, db)
    if pagina < 1 or por_pagina < 10 or por_pagina > POR_PAGINA_MAX:
        raise HTTPException(status_code=422, detail=f"Página inválida (por página: 10 a {POR_PAGINA_MAX}).")
    agora = datetime.utcnow()
    linhas, d1, d2 = _selecionar_lista(db, grupo, situacao, dia, conta, dias, de, ate, galpao, agora, sku)
    total = len(linhas)
    paginas = max((total + por_pagina - 1) // por_pagina, 1)
    pagina = min(pagina, paginas)
    ini = (pagina - 1) * por_pagina
    return {
        "grupo": grupo, "titulo": GRUPOS_LISTA[grupo],
        "usa_periodo": d1 is not None,
        "periodo": {"de": d1.isoformat(), "ate": d2.isoformat()} if d1 else None,
        "filtros": {"situacao": situacao, "dia": dia, "conta": conta, "sku": sku},
        "total": total, "pagina": pagina, "paginas": paginas, "por_pagina": por_pagina,
        "resumo": resumo_da_lista(linhas, agora),
        "itens": [_item_lista(s, agora) for s in linhas[ini:ini + por_pagina]],
    }


@router.get("/lista/excel")
def lista_excel(request: Request, grupo: str = "recebidos", situacao: Optional[str] = None, dia: Optional[str] = None,
                conta: Optional[str] = None, dias: Optional[int] = None, de: Optional[str] = None,
                ate: Optional[str] = None, galpao: Optional[int] = None, sku: Optional[str] = None,
                db: Session = Depends(get_db)):
    """Excel só do grupo clicado (mesmos filtros da lista, sem paginação)."""
    usuario = usuario_gestor_reversos(request, db)
    agora = datetime.utcnow()
    linhas, d1, d2 = _selecionar_lista(db, grupo, situacao, dia, conta, dias, de, ate, galpao, agora, sku)
    r = resumo_da_lista(linhas, agora)

    wb = Workbook()
    aba = wb.active
    aba.title = "Resumo"
    filtros = [x for x in (f"situação: {_nome_situacao(situacao)}" if situacao else "",
                           f"dia: {date.fromisoformat(dia).strftime('%d/%m/%Y')}" if dia else "",
                           f"conta: {conta}" if conta else "",
                           f"SKU: {'sem SKU informado' if sku == SEM_SKU else sku}" if sku else "") if x]
    for linha in (
        ["Lista de reversos"], ["Grupo", GRUPOS_LISTA[grupo]],
        ["Período", f"{d1.strftime('%d/%m/%Y')} a {d2.strftime('%d/%m/%Y')}" if d1 else "Fila de hoje (qualquer data)"],
        ["Galpão", GALPOES[galpao] if galpao else "Todos os galpões"],
        ["Filtros", "; ".join(filtros) or "nenhum"],
        ["Gerado em", f"{datetime.now(FUSO_BR).strftime('%d/%m/%Y %H:%M')} por {usuario.nome_exibicao}"], [],
        ["Total na lista", r["total"]], ["Voltaram ao estoque", r["voltaram"]],
        ["Não voltaram ao estoque", r["nao_voltaram"]], ["Aguardando a equipe", r["aguardando"]],
        ["Tempo médio até decidir", formatar_tempo(r["tempo_medio_min"])], [],
        ["POR SITUAÇÃO / MOTIVO"], *([x["nome"], x["quantidade"]] for x in r["por_motivo"]), [],
        ["POR CONTA", "Reversos", "No estoque", "Fora", "Aguardando"],
    ):
        aba.append(linha)
    # Excel traz TODAS as contas (a tela mostra só as primeiras)
    por_conta = defaultdict(list)
    nomes = {}
    for s in linhas:
        k = s.conta_chave or (s.conta or "").lower()
        nomes.setdefault(k, s.conta or "—")
        por_conta[k].append(s)
    for k, rs in sorted(por_conta.items(), key=lambda kv: (-len(kv[1]), nomes[kv[0]].lower())):
        c = Counter(_situacao_reverso(s) for s in rs)
        v, a = c.get("voltou", 0), c.get("aguardando", 0)
        aba.append([nomes[k], len(rs), v, len(rs) - v - a, a])
    aba.column_dimensions["A"].width = 40
    for col in "BCDE":
        aba.column_dimensions[col].width = 16

    aba_l = wb.create_sheet("Lista")
    aba_l.append(["Recebido em", "Galpão", "Conta", "Plataforma", "Nº da venda", "Produto", "SKU", "Aspecto / motivo relatado",
                  "Situação", "Observação da equipe", "Concluído por", "Concluído em", "Tempo até decidir",
                  "Parado há", "Registrado por", "Origem", "Protocolo"])
    for it in (_item_lista(s, agora) for s in linhas):
        aba_l.append([it["recebido_em"], it["galpao_nome"], it["conta"], it["plataforma"], it["numero_venda"], it["produto"],
                      it["sku"], it["aspecto"], it["situacao_nome"], it["observacao"], it["concluido_por"], it["concluido_em"],
                      formatar_tempo(it["espera_min"]) if it["espera_min"] is not None else "",
                      formatar_tempo(it["parado_ha_min"]) if it["parado_ha_min"] is not None else "",
                      it["registrado_por"], it["origem"], it["protocolo"]])
    for i, larg in enumerate((17, 12, 22, 14, 20, 36, 16, 30, 32, 30, 18, 17, 14, 12, 18, 10, 18), start=1):
        aba_l.column_dimensions[aba_l.cell(row=1, column=i).column_letter].width = larg

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    nome = f"reversos_{grupo}_{datetime.now(FUSO_BR).strftime('%Y-%m-%d')}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{nome}"'},
    )


# ---------------------------------------------------------------------------
# GARANTIAS (27/10): os reversos que a equipe marcou como "foi para garantia", por conta e
# período, já com o CUSTO de cada produto (cadastro de custos da própria conta: CustoSku),
# para a equipe tirar a nota e fazer o procedimento. Só admin/supervisor, sem fotos.
#   GET /api/reversos-gestor/garantias        ?conta=&dias=|&de=&ate=&galpao=  -> JSON
#   GET /api/reversos-gestor/garantias/excel  (mesmos filtros)                 -> Excel
# Custo: casa pela conta + SKU (maiúsculo). Sem custo cadastrado = "sem custo" (nunca inventa
# valor); o total soma só os que têm custo e informa quantos ficaram de fora.
# ---------------------------------------------------------------------------
MAX_ITENS_GARANTIA_TELA = 300   # a tela mostra os mais recentes; o Excel traz todos


def _custos_dos_reversos(db: Session, linhas) -> dict:
    """{(conta_chave, SKU): custo} só para as contas e SKUs que aparecem nos reversos."""
    from app.contas_util import achar_conta_por_nome
    ids_conta, skus = {}, set()
    for s in linhas:
        chave = s.conta_chave or (s.conta or "").lower()
        if chave not in ids_conta:
            conta = achar_conta_por_nome(db, chave)
            ids_conta[chave] = conta.id if conta else None
        skus.update(_skus_do_reverso(s))
    validos = {i for i in ids_conta.values() if i}
    if not validos or not skus:
        return {}
    por_id = {i: k for k, i in ids_conta.items() if i}
    achados = {}
    for c in (db.query(CustoSku).filter(CustoSku.conta_id.in_(validos), func.upper(CustoSku.sku).in_(skus)).all()):
        if c.custo is not None:
            achados[(por_id[c.conta_id], (c.sku or "").strip().upper())] = float(c.custo)
    return achados


def _skus_do_reverso(s: SolicitacaoCancelamento) -> list[str]:
    chave = _chave_sku(s)
    return [] if chave == SEM_SKU else [x.strip() for x in chave.split(",") if x.strip()]


def _item_garantia(s: SolicitacaoCancelamento, custos: dict) -> dict:
    chave = s.conta_chave or (s.conta or "").lower()
    skus = _skus_do_reverso(s)
    valores = [custos.get((chave, k)) for k in skus]
    completo = bool(skus) and all(v is not None for v in valores)
    return {
        "id": s.id,
        "dia": _dia_br(s).isoformat() if s.criado_em else None,
        "recebido_em": _hora_br(s.criado_em),
        "conta": s.conta or "—",
        "plataforma": PLATAFORMAS.get(s.plataforma, s.plataforma or ""),
        "numero_venda": s.numero_venda or "",
        "sku": _chave_sku(s) if skus else "",
        "produto": ((s.produto_titulo or "") if (s.produto_titulo or "").strip().upper() != _chave_sku(s) else ""),
        "custo": round(sum(valores), 2) if completo else None,
        "marcado_por": s.confirmado_por or "",
        "marcado_em": _hora_br(s.confirmado_em),
        "recebido_por": s.solicitado_por or "",
        "observacao": s.observacao_resultado or "",
        "protocolo": _texto_protocolo(s),
    }


def _garantias_selecionadas(db: Session, conta: Optional[str], dias, de, ate, galpao):
    _checar_galpao(galpao)
    d1, d2 = _intervalo(dias, de, ate)
    todas = [s for s in _reversos_do_periodo(db, d1, d2, galpao) if _situacao_reverso(s) == GARANTIA]
    # Lista de contas (com quantas garantias cada uma) NÃO depende do filtro de conta escolhido
    por_conta = Counter(s.conta or "—" for s in todas)
    contas = [{"conta": n, "quantidade": q} for n, q in sorted(por_conta.items(), key=lambda kv: kv[0].lower())]
    linhas = todas
    if conta:
        q = conta.strip().lower()
        linhas = [s for s in todas if (s.conta_chave or "").lower() == q or (s.conta or "").lower() == q]
    linhas = sorted(linhas, key=lambda s: s.criado_em or datetime.min, reverse=True)
    return linhas, contas, d1, d2


def montar_garantias(linhas, custos: dict) -> dict:
    """KPIs e itens das garantias (função pura)."""
    itens = [_item_garantia(s, custos) for s in linhas]
    com_custo = [i["custo"] for i in itens if i["custo"] is not None]
    total = round(sum(com_custo), 2)
    return {
        "kpis": {"quantidade": len(itens), "custo_total": total, "com_custo": len(com_custo),
                 "sem_custo": len(itens) - len(com_custo),
                 "custo_medio": round(total / len(com_custo), 2) if com_custo else None},
        "itens": itens,
    }


@router.get("/garantias")
def garantias(request: Request, conta: Optional[str] = None, dias: Optional[int] = None, de: Optional[str] = None,
              ate: Optional[str] = None, galpao: Optional[int] = None, db: Session = Depends(get_db)):
    usuario_gestor_reversos(request, db)
    linhas, contas, d1, d2 = _garantias_selecionadas(db, conta, dias, de, ate, galpao)
    r = montar_garantias(linhas, _custos_dos_reversos(db, linhas))
    r["total_itens"] = len(r["itens"])
    r["itens"] = r["itens"][:MAX_ITENS_GARANTIA_TELA]
    return {"periodo": {"de": d1.isoformat(), "ate": d2.isoformat()}, "conta": conta or None, "contas": contas, **r}


@router.get("/garantias/excel")
def garantias_excel(request: Request, conta: Optional[str] = None, dias: Optional[int] = None, de: Optional[str] = None,
                    ate: Optional[str] = None, galpao: Optional[int] = None, db: Session = Depends(get_db)):
    usuario = usuario_gestor_reversos(request, db)
    linhas, _contas, d1, d2 = _garantias_selecionadas(db, conta, dias, de, ate, galpao)
    r = montar_garantias(linhas, _custos_dos_reversos(db, linhas))
    k = r["kpis"]
    wb = Workbook()
    aba = wb.active
    aba.title = "Resumo"
    for linha in (
        ["Garantias de reversos"], ["Conta", conta or "Todas as contas"],
        ["Período", f"{d1.strftime('%d/%m/%Y')} a {d2.strftime('%d/%m/%Y')}"],
        ["Galpão", GALPOES[galpao] if galpao else "Todos os galpões"],
        ["Gerado em", f"{datetime.now(FUSO_BR).strftime('%d/%m/%Y %H:%M')} por {usuario.nome_exibicao}"], [],
        ["Reversos em garantia", k["quantidade"]], ["Custo total (só os que têm custo)", k["custo_total"]],
        ["Sem custo cadastrado", k["sem_custo"]], ["Custo médio por garantia", k["custo_medio"] if k["custo_medio"] is not None else ""],
    ):
        aba.append(linha)
    aba.column_dimensions["A"].width = 38
    aba.column_dimensions["B"].width = 26

    aba_g = wb.create_sheet("Garantias")
    aba_g.append(["Recebido em", "Conta", "Plataforma", "Nº da venda", "SKU", "Produto", "Custo (R$)", "Marcado por",
                  "Marcado em", "Recebido no galpão por", "Observação", "Protocolo"])
    for i in r["itens"]:
        aba_g.append([i["recebido_em"], i["conta"], i["plataforma"], i["numero_venda"], i["sku"], i["produto"],
                      i["custo"] if i["custo"] is not None else "sem custo", i["marcado_por"], i["marcado_em"],
                      i["recebido_por"], i["observacao"], i["protocolo"]])
    for n, larg in enumerate((17, 22, 14, 20, 18, 36, 12, 18, 17, 20, 34, 18), start=1):
        aba_g.column_dimensions[aba_g.cell(row=1, column=n).column_letter].width = larg

    aba_s = wb.create_sheet("Por SKU")   # resumo para a nota: quantas unidades de cada SKU, por conta
    aba_s.append(["Conta", "SKU", "Quantidade", "Custo unitário (R$)", "Custo total (R$)"])
    grupos = defaultdict(list)
    for i in r["itens"]:
        grupos[(i["conta"], i["sku"] or "(sem SKU)")].append(i)
    for (cta, sku), its in sorted(grupos.items(), key=lambda kv: (kv[0][0].lower(), kv[0][1])):
        valores = [x["custo"] for x in its if x["custo"] is not None]
        unit = round(valores[0], 2) if valores and len(set(valores)) == 1 else ""
        aba_s.append([cta, sku, len(its), unit if unit != "" else ("sem custo" if not valores else "varia"),
                      round(sum(valores), 2) if valores else "sem custo"])
    for col, larg in zip("ABCDE", (24, 20, 12, 20, 18)):
        aba_s.column_dimensions[col].width = larg

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    nome = f"garantias_{(conta or 'todas').replace(' ', '_')}_{d1.strftime('%Y%m%d')}_{d2.strftime('%Y%m%d')}.xlsx"
    return StreamingResponse(
        buffer, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{nome}"'})
