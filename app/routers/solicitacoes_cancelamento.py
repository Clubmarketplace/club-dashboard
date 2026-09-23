"""
Endpoints das SOLICITAÇÕES MANUAIS de cancelamento -- diferente da fila
de IA em routers/cancelamentos.py (que detecta pedido de cancelamento
dentro de mensagens de clientes). Aqui alguém registra manualmente
"preciso cancelar essa venda" (CEP errado, sem estoque, etiqueta não
gerada...).

Quem pode registrar (a ORIGEM é decidida pelo servidor, pela sessão --
nunca pelo que vem do formulário):
  - "seller":    seller logado; a conta é SEMPRE a vinculada a ele.
  - "logistica": operador de galpão logado; escolhe a conta e o galpão.
  - "publico":   link sem login (continua funcionando como antes).

A data é sempre gerada pelo servidor (datetime.utcnow), nunca vem do
formulário. Um envio pode ter VÁRIAS vendas da mesma conta/plataforma;
cada uma vira uma linha própria (tudo ou nada).

Nome de conta: tudo que compara/agrupa/filtra usa a conta_chave (ver
app/contas_util.py), pra "Friaça", "friaca" e "FRIAÇA" serem a mesma.
"""
import re
import io
from datetime import datetime, timedelta, timezone
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, Request
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from pydantic import BaseModel, field_validator
from sqlalchemy import or_
from sqlalchemy.orm import Session

from app import auth
from app.contas_util import chave_conta, limpar_espacos, nome_exibicao_novo, sufixo_de_plataforma
from app.database import get_db
from app.models import Conta, SolicitacaoCancelamento, Usuario

router = APIRouter(prefix="/api/solicitacoes-cancelamento", tags=["solicitacoes-cancelamento"])

# --- Opções (fonte ÚNICA: formulário, filtros e telas leem daqui) ---
# Pra incluir/renomear: é só mexer nestes dicionários. A chave é o que
# fica gravado no banco -- NUNCA mudar uma chave existente, só o nome.
PLATAFORMAS = {
    "mercado_livre": "Mercado Livre",
    "shopee": "Shopee",
    "magalu": "Magalu",
    "tiktok_shop": "TikTok Shop",
}
GALPOES = {
    1: "Galpão 1",
    2: "Galpão 2",
    3: "Galpão 3",
}
PLATAFORMAS_VALIDAS = set(PLATAFORMAS)
GALPOES_VALIDOS = set(GALPOES)
ORIGENS_VALIDAS = {"seller", "logistica", "publico"}

# Horário de Brasília pros filtros de período. O Brasil não tem horário
# de verão desde 2019, então UTC-3 fixo é exato e não depende de tzdata.
FUSO_BR = timezone(timedelta(hours=-3))

# Papéis que podem CONFIRMAR (quem pede não confirma: controle cruzado).
PAPEIS_QUE_CONFIRMAM = ("admin", "supervisor", "atendente")

RESULTADOS_IMPACTO_VALIDOS = {"sem_impacto", "com_impacto", "aguardando_confirmacao"}
LABEL_RESULTADO_IMPACTO = {
    "sem_impacto": "Sem impacto",
    "com_impacto": "Com impacto",
    "aguardando_confirmacao": "Aguardando confirmação do ML",
}


# ---------------------------------------------------------------------------
# Modelos de entrada
# ---------------------------------------------------------------------------
class ItemCancelamento(BaseModel):
    numero_venda: str
    motivo: str

    @field_validator("numero_venda", "motivo")
    @classmethod
    def validar_nao_vazio(cls, valor: str) -> str:
        valor = valor.strip()
        if not valor:
            raise ValueError("Campo obrigatório não pode ficar em branco.")
        return valor


class NovaSolicitacaoLote(BaseModel):
    plataforma: str
    conta: str
    itens: list[ItemCancelamento]
    galpao: Optional[int] = None  # obrigatório só pra logística
    # Quando o formulário já avisou "essa venda já foi solicitada" e a
    # pessoa decidiu enviar mesmo assim.
    confirmar_duplicadas: bool = False

    @field_validator("plataforma")
    @classmethod
    def validar_plataforma(cls, valor: str) -> str:
        if valor not in PLATAFORMAS_VALIDAS:
            raise ValueError(f"Plataforma inválida. Use uma de: {', '.join(PLATAFORMAS_VALIDAS)}")
        return valor

    @field_validator("conta")
    @classmethod
    def validar_conta(cls, valor: str) -> str:
        valor = limpar_espacos(valor)
        if not valor:
            raise ValueError("Campo obrigatório não pode ficar em branco.")
        return valor

    @field_validator("itens")
    @classmethod
    def validar_ao_menos_um_item(cls, valor: list) -> list:
        if not valor:
            raise ValueError("Informe pelo menos uma venda.")
        return valor


# ---------------------------------------------------------------------------
# Auxiliares
# ---------------------------------------------------------------------------
def _normalizar_numero_venda(numero: str, plataforma: str, estrito: bool) -> str:
    """
    Tira espaços/pontos/traços. No modo estrito (logística): ML só aceita
    números; Shopee, Magalu e TikTok Shop aceitam letras e números, sempre
    em maiúsculo (formato dos pedidos não confirmado pra todas). Fora do
    modo estrito (seller/link público) só limpa espaços -- mantém o
    comportamento que já existia.
    """
    if not estrito:
        return re.sub(r"\s+", "", numero)
    limpo = re.sub(r"[\s.\-/]", "", numero).upper()
    if plataforma == "mercado_livre" and not limpo.isdigit():
        raise HTTPException(status_code=400, detail=f"Nº da venda do Mercado Livre deve ter só números: \"{numero}\".")
    if plataforma != "mercado_livre" and not limpo.isalnum():
        nome = PLATAFORMAS.get(plataforma, plataforma)
        raise HTTPException(status_code=400, detail=f"Nº da venda da {nome} deve ter só letras e números: \"{numero}\".")
    return limpo


def _mapa_contas_conhecidas(db: Session) -> dict[str, str]:
    """
    {chave: nome oficial} de todas as contas que o sistema conhece. Ordem
    de preferência do nome: contas conectadas do ML > conta vinculada dos
    sellers > nome usado em solicitações (o mais antigo, que costuma ser
    o original). Cobre também contas da Shopee (ainda não integradas).
    """
    mapa: dict[str, str] = {}

    def registrar(nome: Optional[str]) -> None:
        if nome and limpar_espacos(nome):
            mapa.setdefault(chave_conta(nome), limpar_espacos(nome))

    for (apelido,) in db.query(Conta.apelido).all():
        registrar(apelido)
    for (vinculada,) in db.query(Usuario.conta_vinculada).filter(Usuario.papel == "seller").all():
        registrar(vinculada)
    for (nome,) in db.query(SolicitacaoCancelamento.conta).order_by(SolicitacaoCancelamento.criado_em.asc()).all():
        registrar(nome)
    return mapa


def _serializar(s: SolicitacaoCancelamento, mapa_contas: Optional[dict] = None) -> dict:
    chave = s.conta_chave or chave_conta(s.conta)
    return {
        "id": s.id,
        "plataforma": s.plataforma,
        "conta": s.conta,
        "conta_chave": chave,
        # Nome oficial da conta (registros antigos podem ter sido digitados
        # como "friaca"; na tela aparece "Friaça").
        "conta_nome": (mapa_contas or {}).get(chave, s.conta),
        "numero_venda": s.numero_venda,
        "motivo": s.motivo,
        "criado_em": s.criado_em.isoformat() if s.criado_em else None,
        "confirmado_por": s.confirmado_por,
        "confirmado_em": s.confirmado_em.isoformat() if s.confirmado_em else None,
        "origem": s.origem,
        "solicitado_por": s.solicitado_por,
        "galpao": s.galpao,
        "galpao_nome": GALPOES.get(s.galpao, f"Galpão {s.galpao}") if s.galpao else None,
        "plataforma_nome": PLATAFORMAS.get(s.plataforma, s.plataforma),
        "resultado_impacto": s.resultado_impacto,
    }


def _inicio_do_dia_br_em_utc(data_br: datetime) -> datetime:
    """Meia-noite de Brasília daquele dia, convertida pra UTC 'cru' (como no banco)."""
    meia_noite = data_br.replace(hour=0, minute=0, second=0, microsecond=0, tzinfo=FUSO_BR)
    return meia_noite.astimezone(timezone.utc).replace(tzinfo=None)


# ---------------------------------------------------------------------------
# Registrar
# ---------------------------------------------------------------------------
@router.post("")
def criar_solicitacoes(corpo: NovaSolicitacaoLote, request: Request, db: Session = Depends(get_db)):
    """
    Registra um ou mais pedidos de uma vez, todos da mesma conta e
    plataforma (tudo ou nada). Rota pública (o link sem login continua
    valendo), mas se houver alguém logado a origem/nome vêm da sessão.
    """
    usuario = auth.usuario_atual(request, db)
    papel = usuario.papel if usuario else None
    mapa_contas = _mapa_contas_conhecidas(db)

    # --- Origem, conta e galpão decididos pelo servidor ---
    galpao = None
    if papel == "seller" and usuario.conta_vinculada:
        origem = "seller"
        conta_digitada = usuario.conta_vinculada  # seller nunca pede por outra conta
    elif papel == "logistica":
        origem = "logistica"
        conta_digitada = corpo.conta
        if corpo.galpao not in GALPOES_VALIDOS:
            raise HTTPException(status_code=400, detail="Escolha o galpão.")
        galpao = corpo.galpao
        sufixo = sufixo_de_plataforma(conta_digitada)
        if sufixo:
            raise HTTPException(
                status_code=400,
                detail=f"Use só o nome da loja, sem \"{sufixo.upper()}\" -- a plataforma já é escolhida nos botões.",
            )
    else:
        origem = "publico"
        conta_digitada = corpo.conta

    chave = chave_conta(conta_digitada)
    # Conta conhecida -> sempre o nome oficial; conta nova -> nome padronizado.
    nome_conta = mapa_contas.get(chave) or nome_exibicao_novo(conta_digitada)

    estrito = origem == "logistica"
    numeros = [_normalizar_numero_venda(i.numero_venda, corpo.plataforma, estrito) for i in corpo.itens]

    # --- Aviso de duplicidade (mesma venda + mesma plataforma) ---
    if not corpo.confirmar_duplicadas:
        existentes = (
            db.query(SolicitacaoCancelamento)
            .filter(SolicitacaoCancelamento.plataforma == corpo.plataforma)
            .filter(SolicitacaoCancelamento.numero_venda.in_(numeros))
            .order_by(SolicitacaoCancelamento.criado_em.asc())
            .all()
        )
        if existentes:
            raise HTTPException(
                status_code=409,
                detail={
                    "codigo": "duplicada",
                    "mensagem": "Já existe solicitação para essa(s) venda(s).",
                    "duplicadas": [
                        {
                            "numero_venda": s.numero_venda,
                            "conta": s.conta,
                            "criado_em": s.criado_em.isoformat() if s.criado_em else None,
                            "solicitado_por": s.solicitado_por,
                            "galpao": s.galpao,
                            "confirmado_por": s.confirmado_por,
                        }
                        for s in existentes
                    ],
                },
            )

    try:
        criadas = []
        for item, numero in zip(corpo.itens, numeros):
            solicitacao = SolicitacaoCancelamento(
                plataforma=corpo.plataforma,
                conta=nome_conta,
                conta_chave=chave,
                numero_venda=numero,
                motivo=item.motivo,
                origem=origem,
                solicitado_por=usuario.nome_exibicao if usuario else None,
                galpao=galpao,
            )
            db.add(solicitacao)
            criadas.append(solicitacao)
        db.commit()
        for s in criadas:
            db.refresh(s)
    except Exception as exc:
        db.rollback()
        raise HTTPException(status_code=500, detail="Não foi possível registrar a solicitação. Tente de novo.") from exc

    return {
        "status": "registrado",
        "quantidade": len(criadas),
        "conta": nome_conta,
        "origem": origem,
        "galpao": galpao,
        "galpao_nome": GALPOES.get(galpao) if galpao else None,
        "solicitado_por": usuario.nome_exibicao if usuario else None,
        "itens": [{"id": s.id, "numero_venda": s.numero_venda, "criado_em": s.criado_em.isoformat()} for s in criadas],
    }


# ---------------------------------------------------------------------------
# Consultar
# ---------------------------------------------------------------------------
@router.get("")
def listar_solicitacoes(db: Session = Depends(get_db), limite: int = 500):
    """Lista simples dos pedidos mais recentes (mantida por compatibilidade)."""
    solicitacoes = (
        db.query(SolicitacaoCancelamento)
        .order_by(SolicitacaoCancelamento.criado_em.desc())
        .limit(min(max(limite, 1), 2000))
        .all()
    )
    return [_serializar(s) for s in solicitacoes]


@router.get("/opcoes")
def listar_opcoes():
    """Plataformas e galpões disponíveis (pro formulário e pros filtros)."""
    return {
        "plataformas": [{"valor": k, "nome": v} for k, v in PLATAFORMAS.items()],
        "galpoes": [{"valor": k, "nome": v} for k, v in GALPOES.items()],
    }


@router.get("/contas")
def listar_contas(db: Session = Depends(get_db)):
    """Lista de contas conhecidas pro campo "Conta" (com busca) e pros filtros."""
    mapa = _mapa_contas_conhecidas(db)
    return sorted(({"chave": k, "nome": v} for k, v in mapa.items()), key=lambda c: c["chave"])


@router.get("/busca")
def buscar_solicitacoes(
    request: Request,
    db: Session = Depends(get_db),
    status: str = "pendente",          # pendente | confirmado | todos
    galpao: Optional[int] = None,
    conta: Optional[str] = None,       # nome ou chave -- comparado pela chave
    plataforma: Optional[str] = None,
    origem: Optional[str] = None,      # seller | logistica | publico
    dias: Optional[int] = 30,          # atalho de período (ignorado se de/ate)
    de: Optional[str] = None,          # AAAA-MM-DD (horário de Brasília)
    ate: Optional[str] = None,
    busca: Optional[str] = None,       # nº da venda ou conta -- ignora o período
    pagina: int = 1,
    por_pagina: int = 100,
):
    """
    Busca com filtros, feita no servidor e paginada. Devolve também os
    contadores (total/pendentes/confirmados) com os mesmos filtros,
    exceto o de status, pros botões mostrarem "Aguardando (3)" etc.
    Logística só enxerga solicitações feitas pelos galpões.
    """
    usuario = auth.usuario_atual(request, db)
    if usuario is None:
        raise HTTPException(status_code=401, detail="Sessão expirada -- faça login de novo.")
    if usuario.papel == "logistica":
        origem = "logistica"

    consulta = db.query(SolicitacaoCancelamento)

    if galpao in GALPOES_VALIDOS:
        consulta = consulta.filter(SolicitacaoCancelamento.galpao == galpao)
    if conta:
        consulta = consulta.filter(SolicitacaoCancelamento.conta_chave == chave_conta(conta))
    if plataforma in PLATAFORMAS_VALIDAS:
        consulta = consulta.filter(SolicitacaoCancelamento.plataforma == plataforma)
    if origem == "publico":
        # Registros antigos (origem nula) entram junto com o link público.
        consulta = consulta.filter(or_(SolicitacaoCancelamento.origem == "publico", SolicitacaoCancelamento.origem.is_(None)))
    elif origem in ORIGENS_VALIDAS:
        consulta = consulta.filter(SolicitacaoCancelamento.origem == origem)

    termo = limpar_espacos(busca or "")
    if termo:
        # Busca geral: em TODO o histórico (ignora o período).
        termo_venda = re.sub(r"\s+", "", termo)
        consulta = consulta.filter(
            or_(
                SolicitacaoCancelamento.numero_venda.ilike(f"%{termo_venda}%"),
                SolicitacaoCancelamento.conta_chave.contains(chave_conta(termo)),
            )
        )
    else:
        agora_br = datetime.now(FUSO_BR)
        try:
            if de or ate:
                if de:
                    inicio = _inicio_do_dia_br_em_utc(datetime.strptime(de, "%Y-%m-%d"))
                    consulta = consulta.filter(SolicitacaoCancelamento.criado_em >= inicio)
                if ate:
                    fim = _inicio_do_dia_br_em_utc(datetime.strptime(ate, "%Y-%m-%d") + timedelta(days=1))
                    consulta = consulta.filter(SolicitacaoCancelamento.criado_em < fim)
            elif dias and dias > 0:
                inicio = _inicio_do_dia_br_em_utc(agora_br - timedelta(days=dias - 1))
                consulta = consulta.filter(SolicitacaoCancelamento.criado_em >= inicio)
        except ValueError:
            raise HTTPException(status_code=400, detail="Data inválida -- use o formato AAAA-MM-DD.")

    # Contadores com todos os filtros, menos o de status.
    total = consulta.count()
    confirmados = consulta.filter(SolicitacaoCancelamento.confirmado_por.isnot(None)).count()

    if status == "pendente":
        consulta = consulta.filter(SolicitacaoCancelamento.confirmado_por.is_(None))
    elif status == "confirmado":
        consulta = consulta.filter(SolicitacaoCancelamento.confirmado_por.isnot(None))

    por_pagina = min(max(por_pagina, 1), 500)
    pagina = max(pagina, 1)
    total_filtrado = consulta.count()
    itens = (
        consulta.order_by(SolicitacaoCancelamento.criado_em.desc())
        .offset((pagina - 1) * por_pagina)
        .limit(por_pagina)
        .all()
    )

    mapa_contas = _mapa_contas_conhecidas(db)
    return {
        "itens": [_serializar(s, mapa_contas) for s in itens],
        "contadores": {"total": total, "confirmados": confirmados, "pendentes": total - confirmados},
        "total_filtrado": total_filtrado,
        "pagina": pagina,
        "por_pagina": por_pagina,
        "tem_mais": pagina * por_pagina < total_filtrado,
        "buscando_todo_historico": bool(termo),
    }


# ---------------------------------------------------------------------------
# Confirmar
# ---------------------------------------------------------------------------
class ConfirmarSolicitacaoBody(BaseModel):
    resultado_impacto: str

    @field_validator("resultado_impacto")
    @classmethod
    def validar_resultado(cls, valor: str) -> str:
        if valor not in RESULTADOS_IMPACTO_VALIDOS:
            raise ValueError(f"resultado_impacto inválido. Use um de: {', '.join(RESULTADOS_IMPACTO_VALIDOS)}")
        return valor


@router.post("/{solicitacao_id}/confirmar")
def confirmar_solicitacao(solicitacao_id: int, corpo: ConfirmarSolicitacaoBody, request: Request, db: Session = Depends(get_db)):
    """
    Marca um pedido como já cancelado de verdade na plataforma. Quem
    confirmou é sempre a pessoa LOGADA no momento do clique (pego da
    sessão, nunca digitado). Seller e logística não confirmam -- quem
    pede não é quem confirma.

    corpo.resultado_impacto é OBRIGATÓRIO e sempre preenchido à mão:
    o Mercado Livre não devolve essa informação por nenhuma API
    pública -- só confirma se pesou ou não na reputação através do
    atendimento deles (chat/WhatsApp, veja os botões de atalho na tela).
    """
    usuario_logado = auth.usuario_atual(request, db)
    if usuario_logado is None:
        raise HTTPException(status_code=401, detail="Sessão expirada -- faça login de novo.")
    if usuario_logado.papel not in PAPEIS_QUE_CONFIRMAM:
        raise HTTPException(status_code=403, detail="Seu perfil não pode confirmar cancelamentos.")

    solicitacao = db.query(SolicitacaoCancelamento).filter(SolicitacaoCancelamento.id == solicitacao_id).first()
    if solicitacao is None:
        raise HTTPException(status_code=404, detail="Solicitação não encontrada")
    if solicitacao.confirmado_por:
        raise HTTPException(status_code=400, detail="Essa solicitação já foi confirmada antes.")

    solicitacao.confirmado_por = usuario_logado.nome_exibicao
    solicitacao.confirmado_em = datetime.utcnow()
    solicitacao.resultado_impacto = corpo.resultado_impacto
    db.commit()
    db.refresh(solicitacao)

    return {
        "status": "confirmado",
        "id": solicitacao.id,
        "confirmado_por": solicitacao.confirmado_por,
        "confirmado_em": solicitacao.confirmado_em.isoformat(),
        "resultado_impacto": solicitacao.resultado_impacto,
    }


@router.post("/{solicitacao_id}/verificar-status")
def verificar_status_agora(solicitacao_id: int, request: Request, db: Session = Depends(get_db)):
    """
    Confere na hora (sem esperar o ciclo automático de até 12 min) se
    essa solicitação pendente já foi cancelada no Mercado Livre --
    útil pra testar uma solicitação específica. Ver
    app/verificacao_cancelamento.py pra lógica completa.
    """
    usuario_logado = auth.usuario_atual(request, db)
    if usuario_logado is None:
        raise HTTPException(status_code=401, detail="Sessão expirada -- faça login de novo.")

    solicitacao = db.query(SolicitacaoCancelamento).filter(SolicitacaoCancelamento.id == solicitacao_id).first()
    if solicitacao is None:
        raise HTTPException(status_code=404, detail="Solicitação não encontrada")
    if solicitacao.confirmado_por:
        raise HTTPException(status_code=400, detail="Essa solicitação já foi confirmada antes.")
    if solicitacao.plataforma != "mercado_livre":
        raise HTTPException(status_code=400, detail="Verificação automática só existe pra Mercado Livre.")

    from app.verificacao_cancelamento import verificar_uma_solicitacao
    confirmou = verificar_uma_solicitacao(solicitacao, db)
    db.refresh(solicitacao)

    return {
        "confirmou": confirmou,
        "mensagem": "Cancelamento confirmado no Mercado Livre!" if confirmou else "Ainda não aparece como cancelada no Mercado Livre.",
        "solicitacao": _serializar(solicitacao),
    }


# ---------------------------------------------------------------------------
# Relatório diário (Excel)
# ---------------------------------------------------------------------------
@router.get("/relatorio")
def relatorio_diario(data: Optional[str] = None, db: Session = Depends(get_db)):
    """
    Gera um Excel (.xlsx) de controle do dia: quantas solicitações
    foram registradas, quantas foram tratadas (confirmadas) e o
    resultado de reputação de cada uma tratada nesse dia -- pra
    acompanhamento diário de quanto foi respondido.

    `data` no formato AAAA-MM-DD (horário de Brasília); se omitido,
    usa o dia de hoje.
    """
    if data:
        try:
            dia = datetime.strptime(data, "%Y-%m-%d")
        except ValueError:
            raise HTTPException(status_code=400, detail="Data inválida -- use o formato AAAA-MM-DD.")
    else:
        dia = datetime.now(FUSO_BR).replace(tzinfo=None)

    inicio_dia = _inicio_do_dia_br_em_utc(dia)
    fim_dia = _inicio_do_dia_br_em_utc(dia + timedelta(days=1))

    tratadas_no_dia = (
        db.query(SolicitacaoCancelamento)
        .filter(SolicitacaoCancelamento.confirmado_em >= inicio_dia, SolicitacaoCancelamento.confirmado_em < fim_dia)
        .order_by(SolicitacaoCancelamento.confirmado_em.asc())
        .all()
    )
    registradas_no_dia = (
        db.query(SolicitacaoCancelamento)
        .filter(SolicitacaoCancelamento.criado_em >= inicio_dia, SolicitacaoCancelamento.criado_em < fim_dia)
        .count()
    )
    pendentes_total = db.query(SolicitacaoCancelamento).filter(SolicitacaoCancelamento.confirmado_por.is_(None)).count()

    contagem_impacto = {"sem_impacto": 0, "com_impacto": 0, "aguardando_confirmacao": 0, "nao_informado": 0}
    for s in tratadas_no_dia:
        chave = s.resultado_impacto if s.resultado_impacto in RESULTADOS_IMPACTO_VALIDOS else "nao_informado"
        contagem_impacto[chave] += 1

    wb = Workbook()

    resumo = wb.active
    resumo.title = "Resumo"
    resumo.append(["Relatório de cancelamentos -- Club Marketplace", dia.strftime("%d/%m/%Y")])
    resumo.append([])
    resumo.append(["Registradas no dia", registradas_no_dia])
    resumo.append(["Tratadas no dia", len(tratadas_no_dia)])
    resumo.append(["Pendentes agora (todas as datas)", pendentes_total])
    resumo.append([])
    resumo.append(["Resultado de reputação das tratadas no dia", ""])
    resumo.append(["Sem impacto", contagem_impacto["sem_impacto"]])
    resumo.append(["Com impacto", contagem_impacto["com_impacto"]])
    resumo.append(["Aguardando confirmação do ML", contagem_impacto["aguardando_confirmacao"]])
    resumo.append(["Não informado", contagem_impacto["nao_informado"]])
    resumo.column_dimensions["A"].width = 40
    resumo.column_dimensions["B"].width = 18

    mapa_contas = _mapa_contas_conhecidas(db)
    detalhe = wb.create_sheet("Detalhe do dia")
    detalhe.append([
        "Conta", "Plataforma", "Nº da venda", "Motivo", "Origem", "Confirmado por",
        "Confirmado em", "Resultado (reputação)",
    ])
    for s in tratadas_no_dia:
        info = _serializar(s, mapa_contas)
        detalhe.append([
            info["conta_nome"],
            info["plataforma_nome"],
            info["numero_venda"],
            info["motivo"],
            {"seller": "Seller", "logistica": info["galpao_nome"] or "Logística", "publico": "Link público"}.get(info["origem"], "Seller / link público"),
            info["confirmado_por"],
            datetime.fromisoformat(info["confirmado_em"]).strftime("%d/%m/%Y %H:%M") if info["confirmado_em"] else "",
            LABEL_RESULTADO_IMPACTO.get(info["resultado_impacto"], "Não informado"),
        ])
    larguras = [18, 14, 18, 30, 18, 18, 18, 26]
    for indice, largura in enumerate(larguras, start=1):
        detalhe.column_dimensions[detalhe.cell(row=1, column=indice).column_letter].width = largura

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    nome_arquivo = f"cancelamentos_{dia.strftime('%Y-%m-%d')}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{nome_arquivo}"'},
    )
