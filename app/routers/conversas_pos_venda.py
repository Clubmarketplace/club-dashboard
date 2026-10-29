"""
Conversas de pós-venda: as mesmas mensagens, em DUAS portas (28/10).

  /api/meu-pos-venda/...  -> SELLER (ou suporte): só as conversas da conta dele.
  /api/pos-venda/conversas... -> EQUIPE (admin, supervisor, atendente): todas as contas.

As duas usam as MESMAS funções (app/pos_venda_conversas.py) e as mesmas linhas do
banco: respondeu numa, some da outra. A conta do seller é sempre a do login -- o
servidor confere em toda rota, nunca a tela.
"""
import io
import os
import threading
from datetime import datetime, timedelta

from fastapi import APIRouter, Depends, HTTPException, Request
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from pydantic import BaseModel
from sqlalchemy.orm import Session

from app import auth
from app import pos_venda_conversas as pv
from app import pos_venda_reclamacoes as pvr
from app.database import get_db
from app.models import MensagemPosVenda

router_seller = APIRouter(prefix="/api/meu-pos-venda", tags=["pós-venda do seller"])
router_equipe = APIRouter(prefix="/api/pos-venda/conversas", tags=["pós-venda da equipe"])


class Resposta(BaseModel):
    texto: str


def _logado(request: Request, db: Session):
    u = auth.usuario_atual(request, db)
    if u is None:
        raise HTTPException(status_code=401, detail="Não autenticado")
    return u


def _usuario_seller(request: Request, db: Session):
    u = _logado(request, db)
    if u.papel not in ("seller", "suporte") or not u.conta_vinculada:
        raise HTTPException(status_code=403, detail="Acesso restrito")
    return u


def _usuario_equipe(request: Request, db: Session):
    u = _logado(request, db)
    if u.papel not in ("admin", "supervisor", "atendente"):
        raise HTTPException(status_code=403, detail="Acesso restrito")
    return u


def _item(db: Session, item_id: int, contas_ids: list[int] | None) -> MensagemPosVenda:
    """A conversa pedida -- e, para o seller, SÓ se for de uma conta dele (senão 404, sem revelar que existe)."""
    q = db.query(MensagemPosVenda).filter(MensagemPosVenda.id == item_id, MensagemPosVenda.tipo == "mensagem")
    if contas_ids is not None:
        q = q.filter(MensagemPosVenda.conta_id.in_(contas_ids or [0]))
    item = q.first()
    if item is None:
        raise HTTPException(status_code=404, detail="Mensagem não encontrada")
    return item


# ----------------------------- funções comuns -----------------------------
def _lista(db: Session, usuario, contas_ids: list[int] | None, conta: str | None = None) -> dict:
    q = db.query(MensagemPosVenda).filter(MensagemPosVenda.tipo == "mensagem", MensagemPosVenda.status.in_((pv.STATUS_ABERTA, "enviando")))
    if contas_ids is not None:
        q = q.filter(MensagemPosVenda.conta_id.in_(contas_ids or [0]))
    abertas = q.order_by(MensagemPosVenda.recebida_em.asc()).all()
    if conta:
        abertas = [i for i in abertas if i.conta and i.conta.apelido == conta]
    desde = datetime.utcnow() - timedelta(days=pv.JANELA_DIAS)
    qr = db.query(MensagemPosVenda).filter(MensagemPosVenda.tipo == "mensagem", MensagemPosVenda.status == "respondida", MensagemPosVenda.respondida_em >= desde)
    if contas_ids is not None:
        qr = qr.filter(MensagemPosVenda.conta_id.in_(contas_ids or [0]))
    respondidas = qr.order_by(MensagemPosVenda.respondida_em.desc()).limit(30).all()
    if conta:
        respondidas = [i for i in respondidas if i.conta and i.conta.apelido == conta]
    return {
        "itens": [pv.para_tela(i, usuario.id) for i in abertas],
        "respondidas": [pv.para_tela(i, usuario.id) for i in respondidas],
        "limite_caracteres": pv.LIMITE_CARACTERES,
        "janela_dias": pv.JANELA_DIAS,
        "inicio": pv._br(pv.inicio_utc()).strftime("%d/%m/%Y"),
        "iniciado": datetime.utcnow() >= pv.inicio_utc(),
    }


def _abrir(db: Session, usuario, item: MensagemPosVenda) -> dict:
    """Abrir = ler a conversa no ML + segurar a trava (se estiver livre) + sugestão da IA."""
    if item.status == "respondida":
        return {"item": pv.para_tela(item, usuario.id), "conversa": _conversa_ou_erro(db, item), "pode_responder": False, "sugestao": None}
    segurou = pv.assumir(db, item.id, usuario)
    db.refresh(item)
    conversa = _conversa_ou_erro(db, item)
    # Alguém já respondeu direto no ML enquanto olhávamos: fecha a linha (a tela mostra "respondida no Mercado Livre").
    if conversa and conversa[-1]["de"] == "vendedor" and item.status == pv.STATUS_ABERTA:
        pv._fechar_externamente(db, item.conta_id, item.pack_id, None)
        db.refresh(item)
    sugestao = pv.sugestao_da_ia(db, item) if segurou and item.status == pv.STATUS_ABERTA else None
    return {"item": pv.para_tela(item, usuario.id), "conversa": conversa, "pode_responder": bool(segurou and item.status == pv.STATUS_ABERTA), "sugestao": sugestao}


def _conversa_ou_erro(db: Session, item: MensagemPosVenda) -> list[dict]:
    try:
        return pv.historico_da_conversa(db, item)
    except pv.RespostaRecusada as exc:
        # Não consegui ler no ML (conta desconectada, ML fora): mostra a última mensagem que temos e o motivo.
        return [{"de": "comprador", "texto": item.texto, "quando": (item.recebida_em.isoformat() + "Z") if item.recebida_em else None}, {"de": "aviso", "texto": exc.mensagem, "quando": None}]


def _responder(db: Session, usuario, item: MensagemPosVenda, texto: str) -> dict:
    try:
        pv.responder(db, item, usuario, texto)
    except pv.RespostaRecusada as exc:
        raise HTTPException(status_code=exc.status, detail=exc.mensagem)
    return {"status": "respondida", "id": item.id}


def _relatorio(db: Session, contas_ids: list[int] | None, dias: int, conta: str | None = None) -> dict:
    # (27/10) Só os últimos 10 dias: o que é mais antigo é descartado do relatório (decisão do Washington).
    dias = pv.JANELA_DIAS
    linhas = pv.linhas_do_relatorio(db, contas_ids, dias, conta)
    return {"dias": dias, "resumo": pv.resumo_do_relatorio(linhas), "itens": [pv.item_do_historico(i) for i in linhas[:500]],
            "reclamacoes": pvr.resumo_do_relatorio(db, contas_ids, conta)}


def _excel(dados: dict, titulo: str, quem: str, com_conta: bool) -> StreamingResponse:
    r = dados["resumo"]
    wb = Workbook()
    a = wb.active
    a.title = "Resumo"
    for linha in (
        [titulo], ["Período", f"últimos {dados['dias']} dias"],
        ["Gerado em", f"{datetime.now(pv.FUSO_BR).strftime('%d/%m/%Y %H:%M')} por {quem}"], [],
        ["Mensagens recebidas", r["recebidas"]],
        ["Respondidas pelo nosso sistema", r["pelo_sistema"]],
        ["   pelo seller", r["pelo_sistema_seller"]], ["   pela equipe", r["pelo_sistema_equipe"]],
        ["Respondidas direto no Mercado Livre", r["direto_no_ml"]],
        ["Aguardando resposta", r["aguardando"]],
        ["% intermediado pelo sistema", r["pct_pelo_sistema"] if r["pct_pelo_sistema"] is not None else ""],
        ["Tempo médio de resposta (min, pelo sistema)", r["tempo_medio_min"] if r["tempo_medio_min"] is not None else ""], [],
        ["POR DIA", "Recebidas", "Pelo sistema"], *([d["dia"], d["recebidas"], d["pelo_sistema"]] for d in r["por_dia"]),
    ):
        a.append(linha)
    if com_conta:
        a.append([])
        a.append(["POR CONTA", "Recebidas", "Pelo sistema", "No ML", "Aguardando"])
        for c in r["por_conta"]:
            a.append([c["conta"], c["recebidas"], c["pelo_sistema"], c["no_ml"], c["aguardando"]])
    a.column_dimensions["A"].width = 46
    for col in "BCDE":
        a.column_dimensions[col].width = 14
    h = wb.create_sheet("Histórico")
    h.append(["Data", "Dia", "Hora", "Conta", "Comprador", "Pedido", "Produto", "Mensagem do comprador", "Situação",
              "Respondida em", "Hora da resposta", "Respondida por", "Como", "Resposta enviada", "Tempo (min)"])
    for i in dados["itens"]:
        h.append([i["recebida"]["data"], i["recebida"]["dia"], i["recebida"]["hora"], i["conta"], i["comprador"], i["pedido"], i["produto"],
                  i["mensagem"], i["status"], i["respondida"]["data"], i["respondida"]["hora"], i["respondida_por"], i["origem"], i["resposta"],
                  i["tempo_min"] if i["tempo_min"] is not None else ""])
    for n, larg in enumerate((11, 9, 7, 18, 18, 18, 34, 50, 18, 12, 9, 18, 24, 50, 10), start=1):
        h.column_dimensions[h.cell(row=1, column=n).column_letter].width = larg
    rec = dados.get("reclamacoes") or {"itens": [], "por_motivo": [], "total": 0}
    c = wb.create_sheet("Reclamações")
    c.append(["Reclamações e mediações (abertas + finalizadas nos últimos 10 dias)", rec["total"]])
    c.append([])
    c.append(["POR MOTIVO", "Total", "Abertas"])
    for m in rec["por_motivo"]:
        c.append([m["motivo"], m["total"], m["abertas"]])
    c.append([])
    c.append(["Atualizada em", "Hora", "Conta", "Comprador", "Pedido", "Produto", "Tipo", "Motivo", "Situação", "Afeta reputação", "Prazo", "Resolução"])
    for i in rec["itens"]:
        c.append([i["atualizada"]["data"], i["atualizada"]["hora"], i["conta"], i["comprador"], i["pedido"], i["produto"], i["tipo"], i["motivo"],
                  i["estado"], i["afeta_texto"], (i["prazo"] or {}).get("data", ""), i["resolucao"]])
    for n, larg in enumerate((34, 8, 18, 18, 18, 34, 14, 30, 28, 24, 12, 30), start=1):
        c.column_dimensions[c.cell(row=1, column=n).column_letter].width = larg
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    nome = f"mensagens_pos_venda_{datetime.now(pv.FUSO_BR).strftime('%Y-%m-%d')}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{nome}"'})


# ------------------------------- SELLER ----------------------------------
@router_seller.get("/relatorio")
def seller_relatorio(request: Request, dias: int = 30, db: Session = Depends(get_db)):
    u = _usuario_seller(request, db)
    return _relatorio(db, pv.contas_ids_do_seller(db, u), dias)


@router_seller.get("/relatorio/excel")
def seller_relatorio_excel(request: Request, dias: int = 30, db: Session = Depends(get_db)):
    u = _usuario_seller(request, db)
    return _excel(_relatorio(db, pv.contas_ids_do_seller(db, u), dias), f"Mensagens pós-venda · {u.conta_vinculada}", u.nome_exibicao, False)


@router_seller.get("/lista")
def seller_lista(request: Request, db: Session = Depends(get_db)):
    u = _usuario_seller(request, db)
    return _lista(db, u, pv.contas_ids_do_seller(db, u))


@router_seller.get("/resumo")
def seller_resumo(request: Request, db: Session = Depends(get_db)):
    u = _usuario_seller(request, db)
    ids = pv.contas_ids_do_seller(db, u)
    n = db.query(MensagemPosVenda).filter(
        MensagemPosVenda.tipo == "mensagem", MensagemPosVenda.status.in_((pv.STATUS_ABERTA, "enviando")), MensagemPosVenda.conta_id.in_(ids or [0])
    ).count()
    return {"aguardando": n, "reclamacoes": pvr.contar_precisam_acao(db, ids)}


@router_seller.get("/reclamacoes")
def seller_reclamacoes(request: Request, db: Session = Depends(get_db)):
    """Reclamações e mediações da conta do seller (só leitura): abertas + finalizadas dos últimos 10 dias."""
    u = _usuario_seller(request, db)
    return pvr.lista_para_tela(db, pv.contas_ids_do_seller(db, u))


@router_seller.post("/{item_id}/abrir")
def seller_abrir(item_id: int, request: Request, db: Session = Depends(get_db)):
    u = _usuario_seller(request, db)
    return _abrir(db, u, _item(db, item_id, pv.contas_ids_do_seller(db, u)))


@router_seller.post("/{item_id}/liberar")
def seller_liberar(item_id: int, request: Request, db: Session = Depends(get_db)):
    u = _usuario_seller(request, db)
    _item(db, item_id, pv.contas_ids_do_seller(db, u))
    pv.liberar(db, item_id, u)
    return {"status": "ok"}


@router_seller.post("/{item_id}/responder")
def seller_responder(item_id: int, corpo: Resposta, request: Request, db: Session = Depends(get_db)):
    u = _usuario_seller(request, db)
    return _responder(db, u, _item(db, item_id, pv.contas_ids_do_seller(db, u)), corpo.texto)


# ------------------------------- EQUIPE ----------------------------------
_varredura = {"rodando": False, "ultimo": None}
_trava_varredura = threading.Lock()


@router_equipe.get("/relatorio")
def equipe_relatorio(request: Request, dias: int = 30, conta: str | None = None, db: Session = Depends(get_db)):
    _usuario_equipe(request, db)
    return _relatorio(db, None, dias, conta)


@router_equipe.get("/relatorio/excel")
def equipe_relatorio_excel(request: Request, dias: int = 30, conta: str | None = None, db: Session = Depends(get_db)):
    u = _usuario_equipe(request, db)
    return _excel(_relatorio(db, None, dias, conta), "Mensagens pós-venda · todas as contas" if not conta else f"Mensagens pós-venda · {conta}", u.nome_exibicao, True)


@router_equipe.get("/lista")
def equipe_lista(request: Request, conta: str | None = None, db: Session = Depends(get_db)):
    u = _usuario_equipe(request, db)
    dados = _lista(db, u, None, conta)
    dados["varredura"] = dict(_varredura)
    dados["contas_reconectar"] = sorted({i["conta"] for i in dados["itens"] if i["conta_reconectar"]})
    return dados


@router_equipe.get("/reclamacoes")
def equipe_reclamacoes(request: Request, conta: str | None = None, db: Session = Depends(get_db)):
    _usuario_equipe(request, db)
    dados = pvr.lista_para_tela(db, None, conta)
    dados["contas_reconectar"] = sorted({i["conta"] for i in dados["itens"] if i["conta_reconectar"]})
    return dados


@router_equipe.post("/{item_id}/abrir")
def equipe_abrir(item_id: int, request: Request, db: Session = Depends(get_db)):
    u = _usuario_equipe(request, db)
    return _abrir(db, u, _item(db, item_id, None))


@router_equipe.post("/{item_id}/liberar")
def equipe_liberar(item_id: int, request: Request, db: Session = Depends(get_db)):
    u = _usuario_equipe(request, db)
    _item(db, item_id, None)
    pv.liberar(db, item_id, u)
    return {"status": "ok"}


@router_equipe.post("/{item_id}/assumir")
def equipe_assumir(item_id: int, request: Request, db: Session = Depends(get_db)):
    """Assumir NO LUGAR de quem está atendendo (só equipe; o seller não toma de ninguém)."""
    u = _usuario_equipe(request, db)
    item = _item(db, item_id, None)
    if not pv.assumir(db, item.id, u, forcar=True):
        raise HTTPException(status_code=409, detail="Essa mensagem já foi respondida.")
    return _abrir(db, u, item)


@router_equipe.post("/{item_id}/responder")
def equipe_responder(item_id: int, corpo: Resposta, request: Request, db: Session = Depends(get_db)):
    u = _usuario_equipe(request, db)
    return _responder(db, u, _item(db, item_id, None), corpo.texto)


@router_equipe.post("/sincronizar")
def equipe_sincronizar(request: Request, anteriores: bool = False, db: Session = Depends(get_db)):
    """
    Busca de novo as conversas de todas as contas (em segundo plano). Normal: só a partir da data de início.
    `anteriores=true` ("Buscar mensagens anteriores"): volta até 10 dias, ignorando a data de início.
    """
    _usuario_equipe(request, db)
    with _trava_varredura:
        if _varredura["rodando"]:
            return {"status": "ja_rodando"}
        _varredura["rodando"] = True

    def _rodar():
        try:
            resultados = pv.sincronizar_todas(anteriores=anteriores)
            try:  # reclamações (29/10): erro aqui não pode estragar a busca das mensagens
                recl = pvr.sincronizar_todas()
            except Exception:
                recl = []
            _varredura["ultimo"] = {
                "em": datetime.utcnow().isoformat() + "Z",
                "criadas": sum(r["criadas"] for r in resultados),
                "reclamacoes": sum(r["criadas"] + r["atualizadas"] for r in recl),
                "contas_reconectar": sorted({r["conta"] for r in resultados + recl if r["erro"] == "reconectar"}),
            }
        finally:
            _varredura["rodando"] = False

    threading.Thread(target=_rodar, daemon=True).start()
    return {"status": "iniciada"}


# ---------------------------------------------------------------------------
# Gestão por conta (28/10) -- visão do supervisor: o que está chegando em cada
# conta, o que está parado e quem respondeu. Só leitura; quem responde continua
# sendo o seller (ou a equipe, nas telas de sempre).
# ---------------------------------------------------------------------------
router_gestao = APIRouter(prefix="/api/pos-venda/gestao", tags=["pós-venda: gestão por conta"])


def _usuario_gestor(request: Request, db: Session):
    u = _logado(request, db)
    if u.papel not in ("admin", "supervisor"):
        raise HTTPException(status_code=403, detail="Acesso restrito")
    return u


def _dias_gestao(dias: int) -> int:
    return max(1, min(int(dias or pv.JANELA_DIAS), pv.JANELA_DIAS))


@router_gestao.get("")
def gestao_resumo(request: Request, dias: int = 10, db: Session = Depends(get_db)):
    """Uma linha por conta liberada: recebidas, aguardando, mais antiga esperando, quem respondeu, tempo médio."""
    from app.models import Conta, ReclamacaoPosVenda
    _usuario_gestor(request, db)
    dias = _dias_gestao(dias)
    agora = datetime.utcnow()

    # Contas (pelo apelido: pode haver duplicatas do mesmo nome) que fazem parte do pós-venda agora.
    contas = [c for c in db.query(Conta).filter(Conta.inativa_em.is_(None)).all() if pv.conta_liberada(c.apelido)]
    por_apelido: dict[str, list] = {}
    for c in contas:
        por_apelido.setdefault(c.apelido, []).append(c)

    linhas = pv.linhas_do_relatorio(db, None, dias)
    por_conta_linhas: dict[str, list] = {}
    for i in linhas:
        por_conta_linhas.setdefault(i.conta.apelido if i.conta else "—", []).append(i)

    # Reclamações abertas que esperam ação, por conta.
    recl = {}
    for r in db.query(ReclamacaoPosVenda).filter(ReclamacaoPosVenda.status_ml == "opened", ReclamacaoPosVenda.precisa_acao.is_(True)).all():
        ap = r.conta.apelido if getattr(r, "conta", None) else None
        if ap:
            recl[ap] = recl.get(ap, 0) + 1

    saida = []
    for apelido in sorted(set(por_apelido) | set(por_conta_linhas), key=str.lower):
        ls = por_conta_linhas.get(apelido, [])
        r = pv.resumo_do_relatorio(ls)
        espera = [(agora - i.recebida_em).total_seconds() / 60 for i in ls if i.status in (pv.STATUS_ABERTA, "enviando") and i.recebida_em]
        grupo_contas = por_apelido.get(apelido, [])
        saida.append({
            "conta": apelido,
            "recebidas": r["recebidas"],
            "aguardando": r["aguardando"],
            "mais_antiga_min": round(max(espera)) if espera else None,
            "pelo_seller": r["pelo_sistema_seller"],
            "pela_equipe": r["pelo_sistema_equipe"],
            "direto_no_ml": r["direto_no_ml"],
            "tempo_medio_min": r["tempo_medio_min"],
            "pct_pelo_sistema": r["pct_pelo_sistema"],
            "reclamacoes_acao": recl.get(apelido, 0),
            "reconectar": any(c.auth_erro_em for c in grupo_contas),
        })
    # Quem tem mais mensagem parada há mais tempo aparece primeiro.
    saida.sort(key=lambda x: (-(x["aguardando"] > 0), -(x["mais_antiga_min"] or 0), x["conta"].lower()))
    total = pv.resumo_do_relatorio(linhas)
    return {
        "dias": dias,
        "contas": saida,
        "totais": {
            "contas": len(saida),
            "recebidas": total["recebidas"],
            "aguardando": total["aguardando"],
            "mais_antiga_min": max((c["mais_antiga_min"] or 0 for c in saida), default=0) or None,
            "reclamacoes_acao": sum(c["reclamacoes_acao"] for c in saida),
            "tempo_medio_min": total["tempo_medio_min"],
            "pct_pelo_sistema": total["pct_pelo_sistema"],
        },
        "verificacao": {"conferir_min": max(2, int(os.getenv("POSVENDA_CONFERIR_MIN", "10") or 10)), "tela_seg": 30},
        "gerado_em": pv.quando_texto(agora),
    }


@router_gestao.get("/mensagens")
def gestao_mensagens(request: Request, conta: str, dias: int = 10, db: Session = Depends(get_db)):
    """As mensagens de UMA conta (o que o comprador disse, se foi respondida, por quem e em quanto tempo)."""
    _usuario_gestor(request, db)
    linhas = pv.linhas_do_relatorio(db, None, _dias_gestao(dias), conta)
    return {"conta": conta, "itens": [pv.item_do_historico(i) for i in linhas[:200]], "total": len(linhas)}


@router_gestao.get("/conversa/{item_id}")
def gestao_conversa(item_id: int, request: Request, db: Session = Depends(get_db)):
    """
    Conversa completa lida do Mercado Livre na hora, SÓ PARA VER: não segura a trava, então o seller
    continua podendo responder. `pode_responder` = ainda está aguardando e ninguém mais está atendendo.
    """
    u = _usuario_gestor(request, db)
    item = _item(db, item_id, None)
    conversa = _conversa_ou_erro(db, item)
    aguardando = item.status == pv.STATUS_ABERTA
    # Alguém respondeu direto no ML: fecha a linha (mesma regra de quando a conversa é aberta nas outras telas).
    if aguardando and conversa and conversa[-1]["de"] == "vendedor":
        pv._fechar_externamente(db, item.conta_id, item.pack_id, None)
        db.refresh(item)
        aguardando = False
    tela = pv.para_tela(item, u.id)
    livre = aguardando and (not tela["em_atendimento_por"] or tela["em_atendimento_eu"])
    return {"item": tela, "conversa": conversa, "aguardando": aguardando, "pode_responder": bool(livre and item.pack_id),
            "limite_caracteres": pv.LIMITE_CARACTERES}


@router_gestao.post("/conversa/{item_id}/responder-agora")
def gestao_habilitar_resposta(item_id: int, request: Request, db: Session = Depends(get_db)):
    """
    O gestor decidiu responder: segura a trava (SEM tomar de quem já está atendendo). Se o seller já estiver
    respondendo, devolve 409 dizendo quem é -- a gestão só acompanha. O envio em si é o /responder de sempre.
    """
    u = _usuario_gestor(request, db)
    item = _item(db, item_id, None)
    if item.status != pv.STATUS_ABERTA:
        raise HTTPException(status_code=409, detail="Essa mensagem já foi respondida.")
    if not pv.assumir(db, item.id, u):
        db.refresh(item)
        raise HTTPException(status_code=409, detail=f"Em atendimento por {item.em_atendimento_por or 'outra pessoa'}. Você só acompanha por enquanto.")
    return {"status": "ok"}
