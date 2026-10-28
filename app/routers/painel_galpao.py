"""
Painel do galpão (26/10) -- números e relatório do que os GALPÕES registraram:
cancelamentos e reversos.

  GET /api/painel-galpao/resumo     ?dias=1|7|15|30 &galpao=   -> JSON (cartões, gráficos, quadros)
  GET /api/painel-galpao/relatorio  ?dias=1|7|15|30 &galpao=   -> Excel (.xlsx)

Regras:
  - Só entra o que foi registrado por galpão (origem "logistica") e só cancelamento
    e reverso. Reputação NUNCA aparece aqui (é assunto do seller/equipe).
  - Quem acessa: logística, admin e supervisor (mesma regra da tela "Acompanhar
    solicitações"). Seller e demais perfis recebem 403.
  - Datas em horário de Brasília; o banco guarda UTC cru.
  - O "dia" de cada registro é o dia em que o galpão registrou (criado_em).
"""
import io
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta, timezone
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, Request
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from sqlalchemy import or_
from sqlalchemy.orm import Session, defer

from app import auth
from app.database import get_db
from app.models import SolicitacaoCancelamento
from app.routers.solicitacoes_cancelamento import (
    AVALIACOES_GALPAO,
    ESTADOS_CHEGADA,
    FUSO_BR,
    GALPOES,
    MOTIVOS_REVERSO,
    PLATAFORMAS,
    RESULTADO_VOLTOU,
    RESULTADOS_REVERSO,
    TIPO_PADRAO,
    TIPO_REVERSO,
    _limites_utc_produtividade,
    _texto_protocolo,
    n_fotos_sql,
)

router = APIRouter(prefix="/api/painel-galpao", tags=["painel-galpao"])

PAPEIS_PAINEL_GALPAO = ("logistica", "admin", "supervisor")
PERIODOS_VALIDOS = (1, 7, 15, 30)   # "Hoje", 7, 15 e 30 dias
TOP_CONTAS = 10


def _usuario_do_painel(request: Request, db: Session):
    usuario = auth.usuario_atual(request, db)
    if usuario is None:
        raise HTTPException(status_code=401, detail="Sessão expirada -- faça login de novo.")
    if usuario.papel not in PAPEIS_PAINEL_GALPAO:
        raise HTTPException(status_code=403, detail="Seu perfil não acessa o painel do galpão.")
    return usuario


def _periodo(dias: int) -> tuple[date, date]:
    if dias not in PERIODOS_VALIDOS:
        raise HTTPException(status_code=422, detail="Período inválido. Use 1 (hoje), 7, 15 ou 30 dias.")
    hoje = datetime.now(FUSO_BR).date()
    return hoje - timedelta(days=dias - 1), hoje


def _registros(db: Session, de: date, ate: date, galpao: Optional[int]):
    """Cancelamentos e reversos registrados pelos galpões no período (sem excluídos)."""
    inicio, fim = _limites_utc_produtividade(de, ate)
    consulta = (
        db.query(SolicitacaoCancelamento)
        # Fotos (base64, até ~800 KB cada) NÃO são lidas: o Dashboard só conta e lista
        # textos. Carregá-las pesava a memória/banco e derrubava o sistema (26/10).
        .options(defer(SolicitacaoCancelamento.fotos_base64))
        .filter(SolicitacaoCancelamento.origem == "logistica")
        .filter(SolicitacaoCancelamento.excluido_em.is_(None))
        .filter(or_(SolicitacaoCancelamento.tipo == TIPO_PADRAO, SolicitacaoCancelamento.tipo == TIPO_REVERSO,
                    SolicitacaoCancelamento.tipo.is_(None)))
        .filter(SolicitacaoCancelamento.criado_em >= inicio, SolicitacaoCancelamento.criado_em < fim)
    )
    if galpao is not None:
        if galpao not in GALPOES:
            raise HTTPException(status_code=422, detail="Galpão inválido.")
        consulta = consulta.filter(SolicitacaoCancelamento.galpao == galpao)
    return consulta.order_by(SolicitacaoCancelamento.criado_em.asc()).all()


def _eh_reverso(s: SolicitacaoCancelamento) -> bool:
    return (s.tipo or TIPO_PADRAO) == TIPO_REVERSO


def _dia_br(s: SolicitacaoCancelamento) -> date:
    return s.criado_em.replace(tzinfo=timezone.utc).astimezone(FUSO_BR).date()


def _situacao_reverso(s: SolicitacaoCancelamento) -> str:
    """Chave de situação do reverso: 'aguardando' | 'voltou' | chave do motivo (não voltou)."""
    if not s.confirmado_por:
        return "aguardando"
    return "voltou" if s.resultado_reverso == RESULTADO_VOLTOU else (s.resultado_reverso or "outro")


def _nome_situacao(chave: str) -> str:
    if chave == "aguardando":
        return "Aguardando a equipe"
    if chave == "voltou":
        return RESULTADOS_REVERSO[RESULTADO_VOLTOU]
    return RESULTADOS_REVERSO.get(chave, "Retornou: sem condição de venda (outro motivo)")


def _tempo_medio_min(reversos) -> Optional[int]:
    """Tempo médio (em minutos) entre o galpão registrar o reverso e a equipe decidir
    (voltou ou não ao estoque). Só entram os já concluídos; sem nenhum, devolve None.
    Defensivo: ignora datas faltando e nunca devolve negativo (relógio/dado ruim)."""
    esperas = []
    for s in reversos:
        if s.confirmado_por and s.confirmado_em and s.criado_em:
            try:
                esperas.append(max((s.confirmado_em - s.criado_em).total_seconds(), 0))
            except TypeError:   # datas com/sem fuso misturadas: ignora esse registro, não derruba o painel
                continue
    return round(sum(esperas) / len(esperas) / 60) if esperas else None


def formatar_tempo(minutos: Optional[int]) -> str:
    """Minutos -> texto curto: '45 min', '3 h 20 min', '2 d 4 h'. None -> '—'."""
    if minutos is None:
        return "—"
    if minutos < 60:
        return f"{minutos} min"
    if minutos < 1440:
        h, m = divmod(minutos, 60)
        return f"{h} h {m:02d} min" if m else f"{h} h"
    d, resto = divmod(minutos, 1440)
    return f"{d} d {resto // 60} h" if resto // 60 else f"{d} d"


def _pct(parte: int, total: int) -> int:
    return round(parte * 100 / total) if total else 0


def montar_resumo(linhas, de: date, ate: date) -> dict:
    """Função pura (testável): números do painel a partir das linhas do período."""
    canc = [s for s in linhas if not _eh_reverso(s)]
    rev = [s for s in linhas if _eh_reverso(s)]

    sit = Counter(_situacao_reverso(s) for s in rev)
    voltaram = sit.get("voltou", 0)
    aguardando = sit.get("aguardando", 0)
    nao_voltaram = len(rev) - voltaram - aguardando
    concluidos = voltaram + nao_voltaram

    # Gráfico por dia (todos os dias do período, mesmo sem registro)
    # Reversos do dia também separados por situação (voltou / não voltou ao estoque / aguardando)
    por_dia = {de + timedelta(days=i): {"cancelamentos": 0, "reversos": 0, "rev_voltou": 0, "rev_nao_voltou": 0, "rev_aguardando": 0}
               for i in range((ate - de).days + 1)}
    for s in linhas:
        d = _dia_br(s)
        if d in por_dia:
            por_dia[d]["reversos" if _eh_reverso(s) else "cancelamentos"] += 1
            if _eh_reverso(s):
                sit_ = _situacao_reverso(s)
                por_dia[d]["rev_voltou" if sit_ == "voltou" else "rev_aguardando" if sit_ == "aguardando" else "rev_nao_voltou"] += 1

    # Resultado dos reversos (voltou, cada motivo de não voltar, aguardando) -- só o que existe
    ordem = ["voltou", "avariado", "faltando_peca", "diferente", "negociacao_ml", "descartado", "outro", "aguardando"]
    extras = [k for k in sit if k not in ordem]       # legados (negociacao_ressarcido etc.)
    resultado = [{"chave": k, "nome": _nome_situacao(k), "quantidade": sit[k]} for k in ordem + extras if sit.get(k)]

    # Quadro por galpão (sempre os 3, pra comparar)
    por_galpao = []
    for numero, nome in GALPOES.items():
        do_galpao = [s for s in linhas if s.galpao == numero]
        c = [s for s in do_galpao if not _eh_reverso(s)]
        r = [s for s in do_galpao if _eh_reverso(s)]
        rs = Counter(_situacao_reverso(s) for s in r)
        por_galpao.append({
            "galpao": numero, "nome": nome,
            "cancelamentos": len(c), "cancelamentos_resolvidos": sum(1 for s in c if s.confirmado_por),
            "reversos": len(r), "voltaram": rs.get("voltou", 0), "aguardando": rs.get("aguardando", 0),
            "nao_voltaram": len(r) - rs.get("voltou", 0) - rs.get("aguardando", 0),
            "pct_voltaram": _pct(rs.get("voltou", 0), len(r) - rs.get("aguardando", 0)),
            "tempo_medio_min": _tempo_medio_min(r),
        })
    # Linha "Total (3 galpões)": soma de todos os galpões listados, pra ninguém somar na mão.
    total_galpoes = {k: sum(g[k] for g in por_galpao)
                     for k in ("cancelamentos", "cancelamentos_resolvidos", "reversos", "voltaram", "nao_voltaram", "aguardando")}
    total_galpoes["pct_voltaram"] = _pct(total_galpoes["voltaram"], total_galpoes["voltaram"] + total_galpoes["nao_voltaram"])
    total_galpoes["tempo_medio_min"] = _tempo_medio_min([s for s in rev if s.galpao in GALPOES])

    # Quem recebeu os reversos no galpão (usuário logado no momento do registro), por galpão.
    quem = defaultdict(int)
    for s in rev:
        quem[(s.solicitado_por or "Não informado", s.galpao)] += 1
    quem_recebeu = sorted(
        ({"nome": n, "galpao": g, "galpao_nome": GALPOES.get(g, "—"), "quantidade": q} for (n, g), q in quem.items()),
        key=lambda x: (-x["quantidade"], x["nome"].lower()),
    )

    # Contas com mais ocorrências
    por_conta = defaultdict(lambda: {"cancelamentos": 0, "reversos": 0})
    nomes = {}
    for s in linhas:
        chave = s.conta_chave or (s.conta or "").lower()
        nomes.setdefault(chave, s.conta)
        por_conta[chave]["reversos" if _eh_reverso(s) else "cancelamentos"] += 1
    contas = sorted(
        ({"conta": nomes[k], **v, "total": v["cancelamentos"] + v["reversos"]} for k, v in por_conta.items()),
        key=lambda x: (-x["total"], x["conta"].lower()),
    )

    return {
        "periodo": {"de": de.isoformat(), "ate": ate.isoformat(), "dias": (ate - de).days + 1},
        "cancelamentos": {
            "total": len(canc),
            "resolvidos": sum(1 for s in canc if s.confirmado_por),
            "pendentes": sum(1 for s in canc if not s.confirmado_por),
        },
        "reversos": {
            "total": len(rev), "voltaram": voltaram, "nao_voltaram": nao_voltaram, "aguardando": aguardando,
            "pct_voltaram": round(voltaram * 100 / concluidos) if concluidos else 0,   # dos já concluídos
            "tempo_medio_min": _tempo_medio_min(rev),
        },
        "por_dia": [{"data": d.isoformat(), **v} for d, v in sorted(por_dia.items())],
        "resultado_reversos": resultado,
        "por_galpao": por_galpao,
        "total_galpoes": total_galpoes,
        "quem_recebeu": quem_recebeu,
        "contas": contas[:TOP_CONTAS],
        "total_contas": len(contas),
    }


@router.get("/resumo")
def resumo(request: Request, dias: int = 7, galpao: Optional[int] = None, db: Session = Depends(get_db)):
    _usuario_do_painel(request, db)
    de, ate = _periodo(dias)
    return montar_resumo(_registros(db, de, ate, galpao), de, ate)


def _hora_br(data_utc: Optional[datetime]) -> str:
    if not data_utc:
        return ""
    return data_utc.replace(tzinfo=timezone.utc).astimezone(FUSO_BR).strftime("%d/%m/%Y %H:%M")


def _situacao_texto(s: SolicitacaoCancelamento) -> str:
    if _eh_reverso(s):
        return _nome_situacao(_situacao_reverso(s))
    return "Resolvido" if s.confirmado_por else "Pendente"


# ---------------------------------------------------------------------------
# Visão explodida (28/10): clicar num cartão, coluna, fatia ou linha abre a LISTA dos registros
# por trás do número, com a situação de cada um e a linha do tempo (recebido -> análise -> decisão).
# Usa as MESMAS linhas do resumo (_registros) e a MESMA regra de situação, então o número do
# cartão e a quantidade da lista são sempre iguais. Só leitura.
# ---------------------------------------------------------------------------
ATRASO_HORAS = 24          # sem decisão da equipe há mais que isso = destaque de atrasado
LIMITE_DETALHE = 500       # teto de linhas por consulta (o Excel da lista traz as mesmas)
_SITUACOES_REVERSO = ("voltou", "nao_voltou", "aguardando", "concluidos")
_SITUACOES_CANCELAMENTO = ("resolvido", "pendente")


def _grupo_da_linha(s: SolicitacaoCancelamento) -> str:
    """Cor/agrupamento da linha: voltou | nao_voltou | aguardando (reversos) | resolvido | pendente (cancelamentos)."""
    if _eh_reverso(s):
        chave = _situacao_reverso(s)
        return "voltou" if chave == "voltou" else "aguardando" if chave == "aguardando" else "nao_voltou"
    return "resolvido" if s.confirmado_por else "pendente"


def _passa_filtro(s, tipo: Optional[str], situacao: Optional[str], dia: Optional[date], conta: Optional[str], quem: Optional[str], q: Optional[str] = None) -> bool:
    if tipo == "reverso" and not _eh_reverso(s):
        return False
    if tipo == "cancelamento" and _eh_reverso(s):
        return False
    if situacao:
        if _eh_reverso(s):
            chave = _situacao_reverso(s)
            grupo = _grupo_da_linha(s)
            if situacao == "concluidos":
                if grupo == "aguardando":
                    return False
            elif situacao in ("voltou", "nao_voltou", "aguardando"):
                if grupo != situacao:
                    return False
            elif chave != situacao:           # chave exata do resultado (avariado, garantia, ...)
                return False
        else:
            if situacao not in _SITUACOES_CANCELAMENTO or _grupo_da_linha(s) != situacao:
                return False
    if dia and _dia_br(s) != dia:
        return False
    if conta and conta.strip().lower() not in ((s.conta or "").lower(), (s.conta_chave or "").lower()):
        return False
    if quem and (s.solicitado_por or "Não informado") != quem:
        return False
    if q and q.strip():   # busca livre: conta, venda, SKU, produto ou quem registrou
        alvo = " ".join(str(x or "") for x in (s.conta, s.numero_venda, s.sku, s.produto_titulo, s.solicitado_por)).lower()
        if q.strip().lower() not in alvo:
            return False
    return True


def _minutos(a: Optional[datetime], b: Optional[datetime]) -> Optional[int]:
    if not a or not b:
        return None
    try:
        return max(round((b - a).total_seconds() / 60), 0)
    except TypeError:     # datas com/sem fuso misturadas: ignora, não derruba a tela
        return None


def _linha_do_tempo(s: SolicitacaoCancelamento, agora: datetime) -> list[dict]:
    """Etapas do registro, na ordem. `feito=False` = etapa que ainda não aconteceu (aguardando)."""
    passos = []
    if _eh_reverso(s):
        det = [x for x in (
            ESTADOS_CHEGADA.get(s.estado_chegada or "", ""), AVALIACOES_GALPAO.get(s.avaliacao_galpao or "", ""),
            MOTIVOS_REVERSO.get(s.motivo_reverso or "", ""), (s.problema_galpao or ""),
        ) if x]
        passos.append({"etapa": "Recebido no galpão", "feito": True, "quando": _hora_br(s.criado_em), "quem": s.solicitado_por or "", "detalhe": " · ".join(det)})
        if s.em_analise_em or s.em_analise_por:
            passos.append({"etapa": "Em análise pela equipe", "feito": True, "quando": _hora_br(s.em_analise_em), "quem": s.em_analise_por or "", "detalhe": s.em_analise_obs or ""})
        elif not s.confirmado_por:
            passos.append({"etapa": "Em análise pela equipe", "feito": False, "quando": "", "quem": s.em_atendimento_por or "", "detalhe": "ainda não pegaram"})
        if s.confirmado_por:
            nome = RESULTADOS_REVERSO.get(s.resultado_reverso or "", _nome_situacao(_situacao_reverso(s)))
            passos.append({"etapa": "Decisão da equipe", "feito": True, "quando": _hora_br(s.confirmado_em or s.liberado_em),
                           "quem": s.confirmado_por or s.liberado_por or "", "detalhe": " · ".join(x for x in (nome, s.observacao_resultado or "") if x)})
        else:
            espera = _minutos(s.criado_em, agora)
            passos.append({"etapa": "Decisão da equipe", "feito": False, "quando": "", "quem": "", "detalhe": f"aguardando há {formatar_tempo(espera)}" if espera is not None else "aguardando"})
    else:
        passos.append({"etapa": "Registrado pelo galpão", "feito": True, "quando": _hora_br(s.criado_em), "quem": s.solicitado_por or "", "detalhe": s.motivo or ""})
        if s.confirmado_por:
            passos.append({"etapa": "Resolvido", "feito": True, "quando": _hora_br(s.confirmado_em), "quem": s.confirmado_por, "detalhe": ("Protocolo: " + _texto_protocolo(s)) if _texto_protocolo(s) else ""})
        else:
            espera = _minutos(s.criado_em, agora)
            passos.append({"etapa": "Resolução", "feito": False, "quando": "", "quem": "", "detalhe": f"pendente há {formatar_tempo(espera)}" if espera is not None else "pendente"})
    return passos


def _item_detalhe(s: SolicitacaoCancelamento, n_fotos: int, agora: datetime) -> dict:
    reverso = _eh_reverso(s)
    decidido = bool(s.confirmado_por)
    espera = None if decidido else _minutos(s.criado_em, agora)
    return {
        "id": s.id,
        "tipo": "reverso" if reverso else "cancelamento",
        "tipo_nome": "Reverso" if reverso else "Cancelamento",
        "grupo": _grupo_da_linha(s),
        "situacao": _situacao_texto(s),
        "situacao_chave": _situacao_reverso(s) if reverso else ("resolvido" if decidido else "pendente"),
        "conta": s.conta,
        "plataforma": s.plataforma,
        "numero_venda": s.numero_venda,
        "sku": s.sku or "",
        "produto": s.produto_titulo or "",
        "galpao": s.galpao,
        "galpao_nome": GALPOES.get(s.galpao, "—"),
        "registrado_por": s.solicitado_por or "",
        "registrado_em": _hora_br(s.criado_em),
        "dia": _dia_br(s).isoformat(),
        "motivo": (MOTIVOS_REVERSO.get(s.motivo_reverso or "", "") if reverso else (s.motivo or "")),
        "espera_min": espera,
        "espera_txt": formatar_tempo(espera) if espera is not None else "",
        "tempo_decisao_min": _minutos(s.criado_em, s.confirmado_em) if decidido else None,
        "tempo_decisao_txt": formatar_tempo(_minutos(s.criado_em, s.confirmado_em)) if decidido and s.confirmado_em else "",
        "atrasado": bool(espera is not None and espera > ATRASO_HORAS * 60),
        "decidido_por": s.confirmado_por or "",
        "n_fotos": int(n_fotos or 0),
        "linha_do_tempo": _linha_do_tempo(s, agora),
    }


def _montar_detalhe(db: Session, dias: int, galpao: Optional[int], tipo, situacao, dia, conta, quem, q=None) -> dict:
    de, ate = _periodo(dias)
    if tipo not in (None, "", "reverso", "cancelamento"):
        raise HTTPException(status_code=422, detail="Tipo inválido.")
    d_dia = None
    if dia:
        try:
            d_dia = date.fromisoformat(dia)
        except ValueError:
            raise HTTPException(status_code=422, detail="Dia inválido (use AAAA-MM-DD).")
    linhas = [s for s in _registros(db, de, ate, galpao) if _passa_filtro(s, tipo or None, situacao or None, d_dia, conta, quem, q)]
    agora = datetime.utcnow()
    pendentes = [s for s in linhas if not s.confirmado_por]
    # Quem espera decisão: o mais antigo primeiro (é o que precisa de atenção). O resto: o mais novo primeiro.
    if linhas and len(pendentes) == len(linhas):
        linhas.sort(key=lambda s: s.criado_em or agora)
    else:
        linhas.sort(key=lambda s: s.criado_em or agora, reverse=True)
    total = len(linhas)
    linhas = linhas[:LIMITE_DETALHE]
    fotos = {}
    if linhas:   # só a QUANTIDADE de fotos (calculada no banco, sem ler foto nenhuma)
        ids = [s.id for s in linhas]
        fotos = {i: n for i, n in db.query(SolicitacaoCancelamento.id, n_fotos_sql()).filter(SolicitacaoCancelamento.id.in_(ids)).all()}
    itens = [_item_detalhe(s, fotos.get(s.id, 0), agora) for s in linhas]
    return {
        "total": total, "truncado": total > len(itens), "limite": LIMITE_DETALHE, "atraso_horas": ATRASO_HORAS,
        "atrasados": sum(1 for i in itens if i["atrasado"]),
        "periodo": {"de": de.isoformat(), "ate": ate.isoformat(), "dias": dias},
        "itens": itens,
    }


@router.get("/detalhe")
def detalhe(request: Request, dias: int = 7, galpao: Optional[int] = None, tipo: Optional[str] = None,
            situacao: Optional[str] = None, dia: Optional[str] = None, conta: Optional[str] = None,
            quem: Optional[str] = None, q: Optional[str] = None, db: Session = Depends(get_db)):
    """Lista dos registros por trás de um cartão/coluna/fatia/linha do Dashboard (mesmos filtros, mesma regra)."""
    _usuario_do_painel(request, db)
    return _montar_detalhe(db, dias, galpao, tipo, situacao, dia, conta, quem, q)


@router.get("/detalhe/excel")
def detalhe_excel(request: Request, dias: int = 7, galpao: Optional[int] = None, tipo: Optional[str] = None,
                  situacao: Optional[str] = None, dia: Optional[str] = None, conta: Optional[str] = None,
                  quem: Optional[str] = None, q: Optional[str] = None, titulo: Optional[str] = None, db: Session = Depends(get_db)):
    """Excel só da lista que o gestor está vendo (uma linha por registro, com a linha do tempo em colunas)."""
    usuario = _usuario_do_painel(request, db)
    d = _montar_detalhe(db, dias, galpao, tipo, situacao, dia, conta, quem, q)
    wb = Workbook()
    aba = wb.active
    aba.title = "Lista"
    aba.append([(titulo or "Registros do galpão")[:120]])
    aba.append([f"Período: {d['periodo']['de']} a {d['periodo']['ate']} · {d['total']} registro(s) · gerado em "
                f"{datetime.now(FUSO_BR).strftime('%d/%m/%Y %H:%M')} por {usuario.nome_exibicao}"])
    aba.append([])
    cab = ["Tipo", "Situação", "Atrasado", "Conta", "Plataforma", "Venda", "SKU", "Produto", "Galpão", "Registrado por", "Registrado em",
           "Motivo", "Espera (sem decisão)", "Tempo até a decisão", "Decidido por", "Fotos", "Linha do tempo"]
    aba.append(cab)
    for i in d["itens"]:
        tempo = " → ".join(f"{p['etapa']}{' (' + p['quando'] + ')' if p['quando'] else ''}{' por ' + p['quem'] if p['quem'] else ''}{': ' + p['detalhe'] if p['detalhe'] else ''}"
                           for p in i["linha_do_tempo"])
        aba.append([i["tipo_nome"], i["situacao"], "SIM" if i["atrasado"] else "", i["conta"], i["plataforma"], i["numero_venda"], i["sku"], i["produto"],
                    i["galpao_nome"], i["registrado_por"], i["registrado_em"], i["motivo"], i["espera_txt"], i["tempo_decisao_txt"], i["decidido_por"], i["n_fotos"], tempo])
    for n, larg in enumerate((12, 38, 9, 18, 12, 20, 14, 34, 11, 18, 17, 24, 18, 18, 18, 7, 90), start=1):
        aba.column_dimensions[aba.cell(row=4, column=n).column_letter].width = larg
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    nome = f"lista_galpao_{datetime.now(FUSO_BR).strftime('%Y-%m-%d_%H%M')}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{nome}"'})


@router.get("/relatorio")
def relatorio(request: Request, dias: int = 7, galpao: Optional[int] = None, db: Session = Depends(get_db)):
    """Excel do período: Resumo, Por conta e Detalhe (uma linha por solicitação)."""
    usuario = _usuario_do_painel(request, db)
    de, ate = _periodo(dias)
    linhas = _registros(db, de, ate, galpao)
    r = montar_resumo(linhas, de, ate)
    rotulo = "Hoje" if dias == 1 else f"Últimos {dias} dias"
    galpao_txt = GALPOES[galpao] if galpao else "Todos os galpões"

    wb = Workbook()
    aba = wb.active
    aba.title = "Resumo"
    for linha in (
        ["Relatório do galpão"], ["Período", f"{rotulo} ({de.strftime('%d/%m/%Y')} a {ate.strftime('%d/%m/%Y')})"],
        ["Galpão", galpao_txt], ["Gerado em", f"{datetime.now(FUSO_BR).strftime('%d/%m/%Y %H:%M')} por {usuario.nome_exibicao}"], [],
        ["CANCELAMENTOS"], ["Registrados", r["cancelamentos"]["total"]], ["Resolvidos", r["cancelamentos"]["resolvidos"]],
        ["Pendentes", r["cancelamentos"]["pendentes"]], [],
        ["REVERSOS"], ["Recebidos", r["reversos"]["total"]], ["No estoque", r["reversos"]["voltaram"]],
        ["Sem condição de venda", r["reversos"]["nao_voltaram"]], ["Aguardando a equipe", r["reversos"]["aguardando"]],
        ["% que voltou ao estoque (dos concluídos)", f"{r['reversos']['pct_voltaram']}%"],
        ["Tempo médio até a equipe decidir", formatar_tempo(r["reversos"]["tempo_medio_min"])], [],
        ["RESULTADO DOS REVERSOS"], *([x["nome"], x["quantidade"]] for x in r["resultado_reversos"]), [],
        ["POR GALPÃO", "Cancelamentos", "Cancel. resolvidos", "Reversos", "No estoque", "Fora do estoque", "Aguardando", "% no estoque", "Tempo médio"],
        *([g["nome"], g["cancelamentos"], g["cancelamentos_resolvidos"], g["reversos"], g["voltaram"], g["nao_voltaram"], g["aguardando"],
           f"{g['pct_voltaram']}%", formatar_tempo(g["tempo_medio_min"])] for g in r["por_galpao"]),
        (lambda g: [f"TOTAL ({galpao_txt})" if galpao else "TOTAL (3 galpões)", g["cancelamentos"], g["cancelamentos_resolvidos"], g["reversos"], g["voltaram"], g["nao_voltaram"],
                    g["aguardando"], f"{g['pct_voltaram']}%", formatar_tempo(g["tempo_medio_min"])])(r["total_galpoes"]),
        [], ["QUEM RECEBEU OS REVERSOS", "Galpão", "Quantidade"],
        *([x["nome"], x["galpao_nome"], x["quantidade"]] for x in r["quem_recebeu"]),
    ):
        aba.append(linha)
    aba.column_dimensions["A"].width = 38
    for col in "BCDEFGHI":
        aba.column_dimensions[col].width = 18

    aba_c = wb.create_sheet("Por conta")
    aba_c.append(["Conta", "Cancelamentos", "Reversos", "Total"])
    # (aqui reagrupa SEM o corte do top 10 que o resumo aplica)
    por_conta = defaultdict(lambda: [0, 0])
    nomes = {}
    for s in linhas:
        k = s.conta_chave or (s.conta or "").lower()
        nomes.setdefault(k, s.conta)
        por_conta[k][1 if _eh_reverso(s) else 0] += 1
    for k, (c, rv) in sorted(por_conta.items(), key=lambda kv: (-(kv[1][0] + kv[1][1]), nomes[kv[0]].lower())):
        aba_c.append([nomes[k], c, rv, c + rv])
    for col, larg in zip("ABCD", (30, 16, 12, 10)):
        aba_c.column_dimensions[col].width = larg

    aba_d = wb.create_sheet("Detalhe")
    aba_d.append(["Registrado em", "Tipo", "Galpão", "Conta", "Plataforma", "Nº da venda", "Produto", "Motivo / aspecto",
                  "Situação", "Concluído por", "Concluído em", "Protocolo", "Observação da equipe", "Registrado por"])
    for s in linhas:
        reverso = _eh_reverso(s)
        aba_d.append([
            _hora_br(s.criado_em), "Reverso" if reverso else "Cancelamento", GALPOES.get(s.galpao, ""), s.conta,
            PLATAFORMAS.get(s.plataforma, s.plataforma), s.numero_venda, s.produto_titulo or s.sku or "", s.motivo or "",
            _situacao_texto(s), s.confirmado_por or "", _hora_br(s.confirmado_em), _texto_protocolo(s),
            (s.observacao_resultado or "") if reverso else "", s.solicitado_por or "",
        ])
    for i, larg in enumerate((17, 14, 11, 22, 14, 20, 36, 30, 30, 18, 17, 18, 30, 18), start=1):
        aba_d.column_dimensions[aba_d.cell(row=1, column=i).column_letter].width = larg

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    nome = f"galpao_{'hoje' if dias == 1 else str(dias) + 'dias'}_{ate.isoformat()}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{nome}"'},
    )
