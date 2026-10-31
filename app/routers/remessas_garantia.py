"""
Remessas de garantia para a fábrica (31/10/2026) -- etapa 2 da Gestão de reversos.

Como a operação funciona: a equipe junta garantias de UMA conta, emite a nota fiscal e envia à fábrica.
O sistema registra esse envio (fábrica, nº da nota, datas, itens, valor) e deixa a remessa "em análise"
até a fábrica responder. Ao responder, o gestor ENCERRA a remessa informando o que voltou (ressarcido em
dinheiro, crédito, troca); o RESTANTE é o prejuízo (`perdido`), calculado aqui no servidor -- nunca digitado.

  POST /api/reversos-gestor/remessas                 cria a remessa (itens = reversos em garantia, de uma conta)
  GET  /api/reversos-gestor/remessas                 lista (?status=analise|encerrada|cancelada|todas &conta= &fabrica= &nota=)
  GET  /api/reversos-gestor/remessas/resumo          cartões do resumo: aguardando agora + recuperado/perdido do período
  GET  /api/reversos-gestor/remessas/excel           planilha (remessas, itens e prejuízo por fábrica)
  GET  /api/reversos-gestor/remessas/{id}            detalhe com itens e histórico
  POST /api/reversos-gestor/remessas/{id}/encerrar   registra a resposta da fábrica (calcula o prejuízo)
  POST /api/reversos-gestor/remessas/{id}/reabrir    SÓ ADMIN: desfaz o encerramento para corrigir
  POST /api/reversos-gestor/remessas/{id}/cancelar   SÓ ADMIN: remessa lançada por engano (libera os reversos)

Regras:
  - Acesso: admin e supervisor (mesma trava da Gestão de reversos). Reabrir e cancelar: só admin.
  - Uma nota = uma conta. Um reverso só entra em UMA remessa ativa (garantido também por índice único no banco).
  - O custo de cada item vem da base de Custos da conta; se faltar, o gestor INFORMA na hora (marcado como
    "informado"); a base de custos nunca é alterada por aqui.
  - Encerrada: valor_total = ressarcido + crédito + troca + perdido (conta fechada em centavos).
  - Tudo que acontece fica em LogRemessaGarantia (só recebe linhas novas). Nada é apagado.
  - Fotos NUNCA são lidas aqui.
"""
import io
from collections import defaultdict
from datetime import date, datetime, timedelta
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, Request
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from pydantic import BaseModel
from sqlalchemy import func
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.database import get_db
from app.models import LogRemessaGarantia, RemessaGarantia, RemessaGarantiaItem, SolicitacaoCancelamento
from app.routers.painel_galpao import _situacao_reverso
from app.routers.reversos_gestor import (
    GARANTIA,
    _base,
    _custos_dos_reversos,
    _intervalo,
    _item_garantia,
    usuario_gestor_reversos,
)
from app.routers.solicitacoes_cancelamento import FUSO_BR, _limites_utc_produtividade

router = APIRouter(prefix="/api/reversos-gestor/remessas", tags=["remessas-garantia"])

STATUS_ANALISE, STATUS_ENCERRADA, STATUS_CANCELADA = "analise", "encerrada", "cancelada"
PRAZO_ATENCAO_DIAS = 15     # passou disso em análise: amarelo ("cobrar a fábrica")
PRAZO_CRITICO_DIAS = 30     # passou disso: vermelho ("esquecida?")
MAX_ITENS_REMESSA = 300
MAX_LISTA = 300


# ---------------------------------------------------------------------------
# Funções puras (testáveis sem banco)
# ---------------------------------------------------------------------------
def _centavos(v) -> int:
    return int(round(float(v) * 100))


def faixa_prazo(dias: int) -> str:
    """'ok' | 'atencao' (>15 dias) | 'critico' (>30 dias)."""
    return "critico" if dias > PRAZO_CRITICO_DIAS else ("atencao" if dias > PRAZO_ATENCAO_DIAS else "ok")


def calcular_encerramento(valor_total: float, ressarcido: float, credito: float, trocado: float) -> dict:
    """
    Fecha a conta de uma remessa. perdido = valor_total - (ressarcido + credito + trocado).
    Levanta ValueError (mensagem em português) se algum valor for negativo ou se o recuperado passar do enviado.
    Tudo em centavos inteiros para não acumular erro de ponto flutuante.
    """
    partes = {"ressarcido": ressarcido, "credito": credito, "trocado": trocado}
    c = {}
    for nome, v in partes.items():
        try:
            c[nome] = _centavos(v or 0)
        except (TypeError, ValueError):
            raise ValueError(f"Valor inválido em {nome}.")
        if c[nome] < 0:
            raise ValueError("Os valores recuperados não podem ser negativos.")
    total = _centavos(valor_total)
    recuperado = c["ressarcido"] + c["credito"] + c["trocado"]
    if recuperado > total:
        raise ValueError(f"O recuperado (R$ {recuperado / 100:.2f}) é maior que o valor enviado (R$ {total / 100:.2f}).")
    return {
        "ressarcido": c["ressarcido"] / 100, "credito": c["credito"] / 100, "trocado": c["trocado"] / 100,
        "recuperado": recuperado / 100, "perdido": (total - recuperado) / 100,
    }


def _hoje_br() -> date:
    return datetime.now(FUSO_BR).date()


def _data(dt: Optional[datetime]) -> Optional[date]:
    return dt.date() if dt else None


def _br(d: Optional[date]) -> str:
    return d.strftime("%d/%m/%Y") if d else ""


def _parse_data(texto: Optional[str], rotulo: str, obrigatoria: bool = True) -> Optional[date]:
    texto = (texto or "").strip()
    if not texto:
        if obrigatoria:
            raise HTTPException(status_code=422, detail=f"Informe a {rotulo}.")
        return None
    try:
        return date.fromisoformat(texto)
    except ValueError:
        raise HTTPException(status_code=422, detail=f"{rotulo.capitalize()} inválida (use o formato AAAA-MM-DD).")


def _dias_em_analise(r: RemessaGarantia, hoje: date) -> int:
    return max(0, (hoje - r.data_envio.date()).days)


def remessa_para_json(r: RemessaGarantia, hoje: date, com_itens: bool = False) -> dict:
    analise = r.status == STATUS_ANALISE
    dias = _dias_em_analise(r, hoje) if analise else None
    d = {
        "id": r.id, "conta": r.conta, "fabrica": r.fabrica, "nota_numero": r.nota_numero,
        "data_emissao": _data(r.data_emissao).isoformat() if r.data_emissao else None,
        "data_envio": _data(r.data_envio).isoformat(), "enviado_por": r.enviado_por or "",
        "observacao_envio": r.observacao_envio or "", "status": r.status,
        "qtd_itens": r.qtd_itens, "valor_total": round(r.valor_total, 2),
        "dias_em_analise": dias, "faixa": faixa_prazo(dias) if analise else None,
        "ressarcido": round(r.ressarcido, 2), "credito": round(r.credito, 2), "trocado": round(r.trocado, 2),
        "recuperado": round(r.ressarcido + r.credito + r.trocado, 2), "perdido": round(r.perdido, 2),
        "data_resposta": _data(r.data_resposta).isoformat() if r.data_resposta else None,
        "respondido_por": r.respondido_por or "", "observacao_resposta": r.observacao_resposta or "",
    }
    if com_itens:
        d["itens"] = [{"id": i.solicitacao_id, "conta": i.conta or "", "numero_venda": i.numero_venda or "",
                       "sku": i.sku or "", "produto": i.produto or "", "custo": round(i.custo, 2),
                       "custo_informado": bool(i.custo_informado)} for i in sorted(r.itens, key=lambda x: x.id)]
    return d


def montar_resumo(remessas_analise, remessas_encerradas, hoje: date) -> dict:
    """Cartões do resumo (função pura): o que aguarda a fábrica agora + o que foi recuperado/perdido no período."""
    aguard = {"remessas": len(remessas_analise), "itens": sum(r.qtd_itens for r in remessas_analise),
              "valor": round(sum(r.valor_total for r in remessas_analise), 2),
              "mais_antiga_dias": max((_dias_em_analise(r, hoje) for r in remessas_analise), default=None),
              "atencao": 0, "criticas": 0}
    for r in remessas_analise:
        f = faixa_prazo(_dias_em_analise(r, hoje))
        aguard["atencao"] += f == "atencao"
        aguard["criticas"] += f == "critico"
    aguard["faixa"] = faixa_prazo(aguard["mais_antiga_dias"]) if aguard["mais_antiga_dias"] is not None else None
    per = {"remessas": len(remessas_encerradas), "enviado": 0.0, "ressarcido": 0.0, "credito": 0.0, "trocado": 0.0,
           "perdido": 0.0}
    fab = defaultdict(lambda: {"remessas": 0, "enviado": 0.0, "recuperado": 0.0, "perdido": 0.0})
    for r in remessas_encerradas:
        per["enviado"] += r.valor_total
        per["ressarcido"] += r.ressarcido
        per["credito"] += r.credito
        per["trocado"] += r.trocado
        per["perdido"] += r.perdido
        f = fab[r.fabrica]
        f["remessas"] += 1
        f["enviado"] += r.valor_total
        f["recuperado"] += r.ressarcido + r.credito + r.trocado
        f["perdido"] += r.perdido
    per = {k: (round(v, 2) if isinstance(v, float) else v) for k, v in per.items()}
    per["recuperado"] = round(per["ressarcido"] + per["credito"] + per["trocado"], 2)
    por_fabrica = sorted(({"fabrica": n, **{k: (round(v, 2) if isinstance(v, float) else v) for k, v in d.items()}}
                          for n, d in fab.items()), key=lambda x: -x["perdido"])
    return {"aguardando": aguard, "periodo": per, "por_fabrica": por_fabrica}


# ---------------------------------------------------------------------------
# Acesso ao banco
# ---------------------------------------------------------------------------
def _filtrar(consulta, conta: Optional[str], fabrica: Optional[str] = None, nota: Optional[str] = None):
    if conta:
        q = conta.strip().lower()
        consulta = consulta.filter((func.lower(RemessaGarantia.conta_chave) == q) | (func.lower(RemessaGarantia.conta) == q))
    if fabrica:
        consulta = consulta.filter(func.lower(RemessaGarantia.fabrica) == fabrica.strip().lower())
    if nota:
        consulta = consulta.filter(RemessaGarantia.nota_numero == nota.strip())
    return consulta


def _carregar(db: Session, remessa_id: int) -> RemessaGarantia:
    r = db.query(RemessaGarantia).filter(RemessaGarantia.id == remessa_id).first()
    if r is None:
        raise HTTPException(status_code=404, detail="Remessa não encontrada.")
    return r


def _log(db: Session, remessa_id: int, acao: str, usuario, detalhe: str) -> None:
    db.add(LogRemessaGarantia(remessa_id=remessa_id, acao=acao, usuario=getattr(usuario, "nome_exibicao", None),
                              detalhe=detalhe))


def _so_admin(usuario) -> None:
    if usuario.papel != "admin":
        raise HTTPException(status_code=403, detail="Só o admin pode fazer isso.")


def _hist(db: Session, remessa_id: int) -> list[dict]:
    linhas = (db.query(LogRemessaGarantia).filter(LogRemessaGarantia.remessa_id == remessa_id)
              .order_by(LogRemessaGarantia.id.asc()).all())
    return [{"acao": l.acao, "usuario": l.usuario or "", "detalhe": l.detalhe or "",
             "quando": l.criado_em.replace(tzinfo=None).isoformat(sep=" ", timespec="minutes")} for l in linhas]


def remessas_dos_reversos(db: Session, ids) -> dict:
    """{id_do_reverso: {remessa_id, nota, fabrica, status, data_envio}} -- usado pela lista de garantias."""
    ids = list(ids)
    if not ids:
        return {}
    linhas = (db.query(RemessaGarantiaItem.solicitacao_id, RemessaGarantia.id, RemessaGarantia.nota_numero,
                       RemessaGarantia.fabrica, RemessaGarantia.status, RemessaGarantia.data_envio)
              .join(RemessaGarantia, RemessaGarantia.id == RemessaGarantiaItem.remessa_id)
              .filter(RemessaGarantiaItem.solicitacao_id.in_(ids)).all())
    return {sid: {"remessa_id": rid, "nota": nota, "fabrica": fab, "status": st, "data_envio": de.date().isoformat()}
            for sid, rid, nota, fab, st, de in linhas}


# ---------------------------------------------------------------------------
# Entradas
# ---------------------------------------------------------------------------
class ItemEntrada(BaseModel):
    id: int
    custo: Optional[float] = None     # só vale quando o custo não está na base de Custos


class RemessaEntrada(BaseModel):
    fabrica: str
    nota_numero: str
    data_envio: str
    data_emissao: Optional[str] = None
    observacao: Optional[str] = None
    itens: list[ItemEntrada]


class EncerrarEntrada(BaseModel):
    ressarcido: float = 0
    credito: float = 0
    trocado: float = 0
    data_resposta: Optional[str] = None
    observacao: Optional[str] = None


# ---------------------------------------------------------------------------
# Rotas (as fixas ANTES das de {id})
# ---------------------------------------------------------------------------
@router.post("")
def criar(dados: RemessaEntrada, request: Request, db: Session = Depends(get_db)):
    usuario = usuario_gestor_reversos(request, db)
    fabrica, nota = (dados.fabrica or "").strip(), (dados.nota_numero or "").strip()
    if not fabrica or len(fabrica) > 80:
        raise HTTPException(status_code=422, detail="Informe a fábrica (até 80 caracteres).")
    if not nota or len(nota) > 40:
        raise HTTPException(status_code=422, detail="Informe o número da nota fiscal (até 40 caracteres).")
    envio = _parse_data(dados.data_envio, "data de envio")
    emissao = _parse_data(dados.data_emissao, "data de emissão", obrigatoria=False)
    if envio > _hoje_br() + timedelta(days=1):
        raise HTTPException(status_code=422, detail="A data de envio não pode estar no futuro.")
    manuais = {i.id: i.custo for i in dados.itens}
    ids = list(dict.fromkeys(i.id for i in dados.itens))
    if not ids:
        raise HTTPException(status_code=422, detail="Escolha pelo menos uma garantia.")
    if len(ids) > MAX_ITENS_REMESSA:
        raise HTTPException(status_code=422, detail=f"No máximo {MAX_ITENS_REMESSA} itens por remessa.")

    linhas = _base(db).filter(SolicitacaoCancelamento.id.in_(ids)).all()
    if len(linhas) != len(ids):
        raise HTTPException(status_code=422, detail="Algum reverso escolhido não existe mais. Atualize a tela.")
    fora = [s for s in linhas if _situacao_reverso(s) != GARANTIA]
    if fora:
        raise HTTPException(status_code=422, detail=f"O reverso da venda {fora[0].numero_venda or fora[0].id} não está em garantia.")
    chaves = {(s.conta_chave or (s.conta or "").lower()) for s in linhas}
    if len(chaves) != 1:
        raise HTTPException(status_code=422, detail="Uma nota é de uma conta só: escolha garantias de uma única conta.")
    ja = remessas_dos_reversos(db, ids)
    if ja:
        sid, info = next(iter(ja.items()))
        raise HTTPException(status_code=409, detail=f"Um dos reversos já está na remessa da nota {info['nota']} ({info['fabrica']}).")

    custos = _custos_dos_reversos(db, linhas)
    itens, faltando = [], []
    for s in sorted(linhas, key=lambda x: x.id):
        it = _item_garantia(s, custos)
        custo, informado = it["custo"], False
        if custo is None:
            manual = manuais.get(s.id)
            if manual is None or manual <= 0:
                faltando.append(s.numero_venda or str(s.id))
                continue
            custo, informado = round(float(manual), 2), True
        itens.append(RemessaGarantiaItem(solicitacao_id=s.id, conta=s.conta, numero_venda=s.numero_venda,
                                         sku=it["sku"], produto=it["produto"], custo=custo, custo_informado=informado))
    if faltando:
        raise HTTPException(status_code=422, detail="Informe o custo de: " + ", ".join(faltando[:10]) + ("…" if len(faltando) > 10 else ""))

    primeira = linhas[0]
    r = RemessaGarantia(
        conta=primeira.conta or "—", conta_chave=next(iter(chaves)), fabrica=fabrica, nota_numero=nota,
        data_emissao=datetime.combine(emissao, datetime.min.time()) if emissao else None,
        data_envio=datetime.combine(envio, datetime.min.time()), observacao_envio=(dados.observacao or "").strip() or None,
        enviado_por=usuario.nome_exibicao, status=STATUS_ANALISE, qtd_itens=len(itens),
        valor_total=round(sum(i.custo for i in itens), 2), itens=itens)
    db.add(r)
    try:
        db.flush()
        _log(db, r.id, "criada", usuario, f"Nota {nota}, fábrica {fabrica}, {len(itens)} itens, R$ {r.valor_total:.2f}")
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=409, detail="Algum reverso acabou de entrar em outra remessa. Atualize a tela.")
    return remessa_para_json(r, _hoje_br(), com_itens=True)


@router.get("")
def listar(request: Request, status: str = STATUS_ANALISE, conta: Optional[str] = None, fabrica: Optional[str] = None,
           nota: Optional[str] = None, db: Session = Depends(get_db)):
    usuario_gestor_reversos(request, db)
    if status not in (STATUS_ANALISE, STATUS_ENCERRADA, STATUS_CANCELADA, "todas"):
        raise HTTPException(status_code=422, detail="Status inválido.")
    consulta = _filtrar(db.query(RemessaGarantia), conta, fabrica, nota)
    if status != "todas":
        consulta = consulta.filter(RemessaGarantia.status == status)
    total = consulta.count()
    linhas = consulta.order_by(RemessaGarantia.data_envio.asc() if status == STATUS_ANALISE else RemessaGarantia.id.desc()) \
        .limit(MAX_LISTA).all()
    hoje = _hoje_br()
    fabricas = [f for (f,) in db.query(RemessaGarantia.fabrica).distinct().order_by(RemessaGarantia.fabrica).all()]
    return {"total": total, "remessas": [remessa_para_json(r, hoje) for r in linhas], "fabricas": fabricas}


@router.get("/resumo")
def resumo(request: Request, conta: Optional[str] = None, dias: Optional[int] = None, de: Optional[str] = None,
           ate: Optional[str] = None, db: Session = Depends(get_db)):
    usuario_gestor_reversos(request, db)
    d1, d2 = _intervalo(dias, de, ate)
    inicio, fim = _limites_utc_produtividade(d1, d2)
    analise = _filtrar(db.query(RemessaGarantia), conta).filter(RemessaGarantia.status == STATUS_ANALISE).all()
    encerradas = (_filtrar(db.query(RemessaGarantia), conta).filter(RemessaGarantia.status == STATUS_ENCERRADA)
                  .filter(RemessaGarantia.encerrada_em >= inicio, RemessaGarantia.encerrada_em < fim).all())
    r = montar_resumo(analise, encerradas, _hoje_br())
    r["periodo_datas"] = {"de": d1.isoformat(), "ate": d2.isoformat()}
    return r


@router.get("/excel")
def excel(request: Request, status: str = "todas", conta: Optional[str] = None, fabrica: Optional[str] = None,
          db: Session = Depends(get_db)):
    usuario = usuario_gestor_reversos(request, db)
    consulta = _filtrar(db.query(RemessaGarantia), conta, fabrica)
    if status in (STATUS_ANALISE, STATUS_ENCERRADA, STATUS_CANCELADA):
        consulta = consulta.filter(RemessaGarantia.status == status)
    remessas = consulta.order_by(RemessaGarantia.id.desc()).all()
    hoje = _hoje_br()
    nomes = {STATUS_ANALISE: "Em análise na fábrica", STATUS_ENCERRADA: "Encerrada", STATUS_CANCELADA: "Cancelada"}
    wb = Workbook()
    aba = wb.active
    aba.title = "Remessas"
    aba.append(["Nota fiscal", "Fábrica", "Conta", "Situação", "Emissão da nota", "Enviada em", "Dias em análise", "Itens",
                "Valor enviado (R$)", "Ressarcido (R$)", "Crédito (R$)", "Trocado (R$)", "Recuperado (R$)", "Perdido (R$)",
                "Respondida em", "Registrada por", "Enviada por", "Obs. do envio", "Obs. da resposta"])
    for r in remessas:
        j = remessa_para_json(r, hoje)
        aba.append([r.nota_numero, r.fabrica, r.conta, nomes.get(r.status, r.status), _br(_data(r.data_emissao)),
                    _br(_data(r.data_envio)), j["dias_em_analise"] if j["dias_em_analise"] is not None else "", r.qtd_itens,
                    j["valor_total"], j["ressarcido"], j["credito"], j["trocado"], j["recuperado"], j["perdido"],
                    _br(_data(r.data_resposta)), r.respondido_por or "", r.enviado_por or "", r.observacao_envio or "",
                    r.observacao_resposta or ""])
    enc = [r for r in remessas if r.status == STATUS_ENCERRADA]
    aba.append(["TOTAL (encerradas)", "", "", "", "", "", "", sum(r.qtd_itens for r in enc), round(sum(r.valor_total for r in enc), 2),
                round(sum(r.ressarcido for r in enc), 2), round(sum(r.credito for r in enc), 2), round(sum(r.trocado for r in enc), 2),
                round(sum(r.ressarcido + r.credito + r.trocado for r in enc), 2), round(sum(r.perdido for r in enc), 2)])
    for n, larg in enumerate((16, 20, 22, 22, 14, 14, 12, 8, 16, 14, 12, 12, 16, 14, 14, 18, 18, 30, 30), start=1):
        aba.column_dimensions[aba.cell(row=1, column=n).column_letter].width = larg

    aba_i = wb.create_sheet("Itens")
    aba_i.append(["Nota fiscal", "Fábrica", "Conta", "Nº da venda", "SKU", "Produto", "Custo (R$)", "Custo informado na hora"])
    for r in remessas:
        for i in sorted(r.itens, key=lambda x: x.id):
            aba_i.append([r.nota_numero, r.fabrica, r.conta, i.numero_venda or "", i.sku or "", i.produto or "",
                          round(i.custo, 2), "sim" if i.custo_informado else ""])
    for n, larg in enumerate((16, 20, 22, 20, 18, 36, 12, 22), start=1):
        aba_i.column_dimensions[aba_i.cell(row=1, column=n).column_letter].width = larg

    aba_f = wb.create_sheet("Prejuízo por fábrica")
    aba_f.append(["Fábrica", "Remessas encerradas", "Enviado (R$)", "Recuperado (R$)", "Perdido (R$)", "% recuperado"])
    res = montar_resumo([], enc, hoje)
    for f in res["por_fabrica"]:
        aba_f.append([f["fabrica"], f["remessas"], f["enviado"], f["recuperado"], f["perdido"],
                      round(f["recuperado"] * 100 / f["enviado"], 1) if f["enviado"] else ""])
    p = res["periodo"]
    aba_f.append(["TOTAL", p["remessas"], p["enviado"], p["recuperado"], p["perdido"],
                  round(p["recuperado"] * 100 / p["enviado"], 1) if p["enviado"] else ""])
    for n, larg in enumerate((24, 20, 16, 16, 16, 14), start=1):
        aba_f.column_dimensions[aba_f.cell(row=1, column=n).column_letter].width = larg
    aba_r = wb.create_sheet("Gerado")
    aba_r.append(["Gerado em", f"{datetime.now(FUSO_BR).strftime('%d/%m/%Y %H:%M')} por {usuario.nome_exibicao}"])
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="remessas_garantia.xlsx"'})


@router.get("/{remessa_id}")
def detalhe(remessa_id: int, request: Request, db: Session = Depends(get_db)):
    usuario_gestor_reversos(request, db)
    r = _carregar(db, remessa_id)
    j = remessa_para_json(r, _hoje_br(), com_itens=True)
    j["historico"] = _hist(db, r.id)
    return j


@router.post("/{remessa_id}/encerrar")
def encerrar(remessa_id: int, dados: EncerrarEntrada, request: Request, db: Session = Depends(get_db)):
    usuario = usuario_gestor_reversos(request, db)
    r = _carregar(db, remessa_id)
    if r.status != STATUS_ANALISE:
        raise HTTPException(status_code=409, detail="Esta remessa não está mais em análise.")
    try:
        c = calcular_encerramento(r.valor_total, dados.ressarcido, dados.credito, dados.trocado)
    except ValueError as e:
        raise HTTPException(status_code=422, detail=str(e))
    resposta = _parse_data(dados.data_resposta, "data da resposta", obrigatoria=False) or _hoje_br()
    if resposta > _hoje_br() + timedelta(days=1):
        raise HTTPException(status_code=422, detail="A data da resposta não pode estar no futuro.")
    if resposta < r.data_envio.date():
        raise HTTPException(status_code=422, detail="A resposta não pode ser anterior à data de envio.")
    r.ressarcido, r.credito, r.trocado, r.perdido = c["ressarcido"], c["credito"], c["trocado"], c["perdido"]
    r.data_resposta = datetime.combine(resposta, datetime.min.time())
    r.respondido_por = usuario.nome_exibicao
    r.observacao_resposta = (dados.observacao or "").strip() or None
    r.encerrada_em = datetime.utcnow()
    r.status = STATUS_ENCERRADA
    _log(db, r.id, "encerrada", usuario,
         f"Enviado R$ {r.valor_total:.2f}; ressarcido R$ {c['ressarcido']:.2f}, crédito R$ {c['credito']:.2f}, "
         f"troca R$ {c['trocado']:.2f}; perdido R$ {c['perdido']:.2f}")
    db.commit()
    return remessa_para_json(r, _hoje_br(), com_itens=True)


@router.post("/{remessa_id}/reabrir")
def reabrir(remessa_id: int, request: Request, db: Session = Depends(get_db)):
    usuario = usuario_gestor_reversos(request, db)
    _so_admin(usuario)
    r = _carregar(db, remessa_id)
    if r.status != STATUS_ENCERRADA:
        raise HTTPException(status_code=409, detail="Só dá para reabrir uma remessa encerrada.")
    antes = (f"ressarcido R$ {r.ressarcido:.2f}, crédito R$ {r.credito:.2f}, troca R$ {r.trocado:.2f}, "
             f"perdido R$ {r.perdido:.2f}, resposta em {_br(_data(r.data_resposta))} por {r.respondido_por or '—'}")
    r.status = STATUS_ANALISE
    r.ressarcido = r.credito = r.trocado = r.perdido = 0.0
    r.data_resposta = r.encerrada_em = None
    r.respondido_por = r.observacao_resposta = None
    _log(db, r.id, "reaberta", usuario, "Valores anteriores: " + antes)
    db.commit()
    return remessa_para_json(r, _hoje_br(), com_itens=True)


@router.post("/{remessa_id}/cancelar")
def cancelar(remessa_id: int, request: Request, db: Session = Depends(get_db)):
    usuario = usuario_gestor_reversos(request, db)
    _so_admin(usuario)
    r = _carregar(db, remessa_id)
    if r.status != STATUS_ANALISE:
        raise HTTPException(status_code=409, detail="Só dá para cancelar uma remessa ainda em análise (reabra a encerrada antes).")
    lista = "; ".join(f"{i.numero_venda or i.solicitacao_id}/{i.sku or '-'}/R$ {i.custo:.2f}" for i in r.itens)
    _log(db, r.id, "cancelada", usuario, f"Itens liberados: {lista}")
    r.itens.clear()        # libera os reversos para entrarem em outra remessa
    r.status = STATUS_CANCELADA
    db.commit()
    return remessa_para_json(r, _hoje_br())
